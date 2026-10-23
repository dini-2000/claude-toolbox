#!/usr/bin/env python3
"""Convert a House Style deck into an editable PowerPoint that matches it 1:1.

    python3 tools/to_pptx.py samples/dist/showcase.html              # writes showcase.pptx beside it
    python3 tools/to_pptx.py deck.html -o out.pptx --check            # and compare every slide
    python3 tools/to_pptx.py deck.html --check --powerpoint           # also check in PowerPoint (Mac)

The method of PDF-to-PowerPoint converters (Acrobat, iLovePDF):
1. Chrome prints the deck with its own print CSS, one exact page per slide at the deck's format (16:9, square,
   portrait, A4 ...), three times: the whole slide, the artwork alone and the content alone. Artwork is the
   slide's background, its ::before and ::after layers, elements with class pc-art and full-bleed pictures
   (pc-media--bg), so a theme's decoration stays exactly as drawn.
2. PyMuPDF reads what Chrome drew on each content page: every piece of text with its position, font, size and
   colour, every vector shape and every image.
3. python-pptx rebuilds each slide:
   - artwork becomes one picture at the back (or a plain background colour);
   - text becomes editable text boxes in the theme's own fonts, with the same line breaks, one paragraph per
     line so PowerPoint, LibreOffice and Google Slides all place each line where Chrome did; each box has room
     for its lines in Arial too, for apps that lack the font;
   - shapes become native rectangles, rounded rectangles, ellipses and freeforms; images, and what PowerPoint
     cannot draw natively (soft shadows, gradients), become pictures;
   - speaker notes come from <aside class="pc-notes">.
   Every font face the slides use is embedded (the kit's fonts are all OFL), and the file's theme gets the
   deck's fonts and colours, so text and shapes added later match the deck.
--check converts the PowerPoint back to PDF with LibreOffice and writes, for every slide, the HTML slide and the
PowerPoint slide side by side plus a difference score and a line-position test.

Needs Google Chrome (or Chromium or Edge), Python 3.9+, PyMuPDF and python-pptx; --check needs LibreOffice.
A deck that links the kit's engine/ files is built into one file first (tools/build.py).
"""
import argparse
import glob
import hashlib
import io
import math
from functools import lru_cache
import os
import re
import shutil
import signal
import subprocess
import sys
import tempfile
import time
import zipfile
from html.parser import HTMLParser
from pathlib import Path

try:
    import fitz  # PyMuPDF
except ImportError:
    sys.exit('PyMuPDF is missing. Install it with:  python3 -m pip install pymupdf')
try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.opc.package import Part
    from pptx.opc.packuri import PackURI
    from pptx.oxml import parse_xml
    from pptx.oxml.ns import qn
except ImportError:
    sys.exit('python-pptx is missing. Install it with:  python3 -m pip install python-pptx')

HERE = Path(__file__).resolve().parent
KIT = HERE.parent
sys.path.insert(0, str(HERE))
import chrome    # noqa: E402
import fontinfo  # noqa: E402

EMU = 12700                      # EMU per PDF point (Chrome prints 1 CSS px as 0.75 pt)
ART_ZOOM = 8 / 3                 # artwork picture at twice the slide's CSS pixels
NS = ('xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
      'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
      'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"')
RT_FONT = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/font'

# Print variants (engine/core.css: "Art layer"). Full-bleed pictures go with the artwork, since a theme's scrim
# (a ::before layer) sits between the picture and the text.
ART = '.pc-slide > .pc-media--bg, .pc-slide .pc-art'
VARIANT_CSS = {
    'full': '',
    'art': ('.pc-slide * { visibility: hidden !important; }'
            '%s, :is(%s) * { visibility: visible !important; }' % (ART, ART)),
    'content': ('.pc-slide::before, .pc-slide::after { display: none !important; }'
                '.pc-slide { background: transparent !important; box-shadow: none !important; }'
                '%s { visibility: hidden !important; }' % ART),
}
COMMON_CSS = '*, *::before, *::after { animation: none !important; transition: none !important; }'
# Shapes drawn with a CSS mask (an arrow, a sticker burst) print as a fill plus the mask's own drawing, which
# PyMuPDF cannot tell apart; a script tags them (data-pc-mask: self, before, after), the content print leaves
# them out and a fourth print holds only them, on a transparent page, to place as pictures.
MASK_JS = ('<script>(function(){function tag(){document.querySelectorAll(".pc-slide *").forEach(function(el){'
           'if(el.closest(".pc-art,.pc-media--bg"))return;var t=[];["","::before","::after"].forEach(function(p){'
           'var cs=getComputedStyle(el,p||null),m=cs.maskImage||cs.webkitMaskImage;'
           'if(m&&m!=="none"&&(!p||(cs.content!=="none"&&cs.content!=="normal")))t.push(p?p.slice(2):"self")});'
           'if(t.length)el.setAttribute("data-pc-mask",t.join(" "))})}'
           'addEventListener("load",tag);addEventListener("beforeprint",tag)})();</script>')
MASKED = '[data-pc-mask~="self"], [data-pc-mask~="before"]::before, [data-pc-mask~="after"]::after'
VARIANT_CSS['content'] += '%s { visibility: hidden !important; }' % MASKED
VARIANT_CSS['masks'] = ('html, body, .pc-deck, .pc-slide { background: transparent !important; box-shadow: none !important; }'
                        '.pc-slide::before, .pc-slide::after { display: none !important; }'
                        '.pc-slide * { visibility: hidden !important; }'
                        '%s, [data-pc-mask~="self"] * { visibility: visible !important; }' % MASKED)


def log(*a):
    print(*a, flush=True)


# ------------------------------------------------------------------ finding the tools

def find_soffice():
    cands = [os.environ.get('SOFFICE_PATH', ''), shutil.which('soffice') or '', shutil.which('libreoffice') or '',
             '/Applications/LibreOffice.app/Contents/MacOS/soffice',
             os.path.expanduser('~/Applications/LibreOffice.app/Contents/MacOS/soffice'),
             r'C:\Program Files\LibreOffice\program\soffice.exe',
             r'C:\Program Files (x86)\LibreOffice\program\soffice.exe']
    return next((c for c in cands if c and os.path.isfile(c)), None)


# ------------------------------------------------------------------ printing

def print_variant(deck, variant, pdf, expect=None):
    """Print one variant of a one-file deck with headless Chrome (tools/chrome.py fires beforeprint once the
    deck is ready, which ends builds and draws charts still undrawn)."""
    style = '<style>%s%s</style>%s' % (COMMON_CSS, VARIANT_CSS[variant], MASK_JS if variant != 'full' else '')
    return chrome.print_pdf(deck, pdf, expect=expect, head=style)


def one_file(src, work):
    """The deck as one file: itself when it is built already, else a build of it in the work folder."""
    text = Path(src).read_text(encoding='utf-8')
    if re.search(r'(?:href|src)=["\'][^"\']*engine/(?:core\.css|runtime\.js)', text):
        import build
        out = Path(work) / (Path(src).stem + '.html')
        build.build(src, out=out, quiet=True, check=False)
        return out
    return Path(src)

# ------------------------------------------------------------------ speaker notes

VOID = {'area', 'base', 'br', 'col', 'embed', 'hr', 'img', 'input', 'link', 'meta', 'source', 'track', 'wbr'}
BLOCK = {'p', 'li', 'div', 'ul', 'ol', 'h1', 'h2', 'h3', 'h4', 'section', 'tr'}


class NotesParser(HTMLParser):
    """Collect the text of each .pc-notes element, per .pc-slide, in document order.
    Skips <noscript>, where tools/build.py keeps a copy of the slides for viewers without scripts."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.slides, self.stack, self.depth, self.noscript = [], [], 0, 0

    def handle_starttag(self, tag, attrs):
        if tag == 'noscript':
            self.noscript += 1
        if self.noscript:
            return
        if tag in VOID:
            if tag == 'br' and self.depth:
                self.slides[-1][-1] += '\n'
            return
        cls = (dict(attrs).get('class') or '').split()
        slide, notes = 'pc-slide' in cls, 'pc-notes' in cls
        self.stack.append((tag, notes))
        if slide:
            self.slides.append([])
        if notes and self.slides:
            self.depth += 1
            if self.depth == 1:
                self.slides[-1].append('')
        elif self.depth and tag in BLOCK:
            self.slides[-1][-1] += '\n'

    def handle_endtag(self, tag):
        if self.noscript:
            if tag == 'noscript':
                self.noscript -= 1
            return
        while self.stack:
            t, notes = self.stack.pop()
            if notes:
                self.depth -= 1
            if t == tag:
                break

    def handle_data(self, data):
        if self.depth and self.slides and not self.noscript:
            self.slides[-1][-1] += re.sub(r'\s+', ' ', data)


def read_notes(html_path):
    p = NotesParser()
    p.feed(Path(html_path).read_text(encoding='utf-8'))
    out = []
    for parts in p.slides:
        text = '\n\n'.join(re.sub(r' *\n *', '\n', s).strip() for s in parts if s.strip())
        out.append(re.sub(r'\n{3,}', '\n\n', text))
    return out


# ------------------------------------------------------------------ XML helpers

def esc(s):
    s = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', s)
    return s.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;').replace('"', '&quot;')


def rgb_hex(color, n=3):
    if color is None:
        return None
    c = list(color)
    if len(c) == 1:
        c = c * 3
    elif len(c) == 4:
        k = c[3]
        c = [(1 - c[i]) * (1 - k) for i in range(3)]
    return '%02X%02X%02X' % tuple(max(0, min(255, int(round(v * 255)))) for v in c[:3])


def solid(hexcol, opacity=1.0):
    alpha = '' if opacity >= 0.999 else '<a:alpha val="%d"/>' % int(round(max(0, opacity) * 100000))
    return '<a:solidFill><a:srgbClr val="%s">%s</a:srgbClr></a:solidFill>' % (hexcol, alpha)


def emu(v):
    return int(round(v * EMU))


class Ids:
    def __init__(self, start):
        self.n = start

    def next(self):
        self.n += 1
        return self.n


# ------------------------------------------------------------------ vector shapes

def style_key(d):
    return (d.get('type'), rgb_hex(d.get('fill')), rgb_hex(d.get('color')), round(d.get('width') or 0, 3),
            round(d.get('fill_opacity') or 0, 3), round(d.get('stroke_opacity') or 0, 3), d.get('even_odd'),
            str(d.get('dashes')), str(d.get('lineCap')), str(d.get('lineJoin')))


def subpaths(d):
    """Split PyMuPDF path items into subpaths of ('M'|'L'|'C', points...)."""
    subs, cur, last = [], None, None
    close = bool(d.get('closePath'))

    def near(p, q):
        return abs(p.x - q.x) < 0.01 and abs(p.y - q.y) < 0.01

    for it in d['items']:
        kind = it[0]
        if kind == 're':
            r = it[1]
            pts = [r.tl, r.tr, r.br, r.bl]
            if len(it) > 2 and it[2] == -1:
                pts = [r.tl, r.bl, r.br, r.tr]
            subs.append({'ops': [('M', pts[0])] + [('L', p) for p in pts[1:]], 'closed': True})
            cur, last = None, None
            continue
        if kind == 'qu':
            q = it[1]
            subs.append({'ops': [('M', q.ul), ('L', q.ur), ('L', q.lr), ('L', q.ll)], 'closed': True})
            cur, last = None, None
            continue
        p0 = it[1]
        if cur is None or last is None or not near(last, p0):
            cur = {'ops': [('M', p0)], 'closed': False}
            subs.append(cur)
        if kind == 'l':
            cur['ops'].append(('L', it[2]))
            last = it[2]
        elif kind == 'c':
            cur['ops'].append(('C', it[2], it[3], it[4]))
            last = it[4]
    for s in subs:
        ops = s['ops']
        endp = ops[-1][-1]
        if not s['closed'] and (close or 'f' in (d.get('type') or '')) and len(ops) > 1:
            s['closed'] = close or near(ops[0][1], endp)
    return subs


def signed_area(sub):
    pts = [op[-1] for op in sub['ops']]
    return sum(pts[i].x * pts[(i + 1) % len(pts)].y - pts[(i + 1) % len(pts)].x * pts[i].y
               for i in range(len(pts))) / 2


def detect_prst(subs):
    """Recognise a plain rectangle, rounded rectangle or ellipse, for nicer editing.
    Chrome draws each rounded corner as several curves, so fit the points instead of counting."""
    if len(subs) != 1:
        return None
    ops = subs[0]['ops']
    on = [o[-1] for o in ops]
    x0, x1 = min(p.x for p in on), max(p.x for p in on)
    y0, y1 = min(p.y for p in on), max(p.y for p in on)
    w, h = x1 - x0, y1 - y0
    if w <= 0.05 or h <= 0.05:
        return None
    kinds = [o[0] for o in ops[1:]]
    if 'C' not in kinds:
        if len(on) in (4, 5) and all(min(abs(p.x - x0), abs(p.x - x1)) < 0.05 and
                                     min(abs(p.y - y0), abs(p.y - y1)) < 0.05 for p in on):
            return ('rect', x0, y0, w, h, None)
        return None
    cx, cy, rx, ry = (x0 + x1) / 2, (y0 + y1) / 2, w / 2, h / 2
    if all(abs(((p.x - cx) / rx) ** 2 + ((p.y - cy) / ry) ** 2 - 1) < 0.03 for p in on):
        return ('ellipse', x0, y0, w, h, None)
    cand = []
    prev = ops[0][1]
    for o in ops[1:]:
        a, b = prev, o[-1]
        if o[0] == 'L':
            if abs(a.x - b.x) < 0.05 and min(abs(a.x - x0), abs(a.x - x1)) < 0.05:
                cand += [min(a.y, b.y) - y0, y1 - max(a.y, b.y)]
            elif abs(a.y - b.y) < 0.05 and min(abs(a.y - y0), abs(a.y - y1)) < 0.05:
                cand += [min(a.x, b.x) - x0, x1 - max(a.x, b.x)]
        prev = b
    cand = sorted(c for c in cand if c > 0.05)
    r = cand[len(cand) // 2] if cand else min(w, h) / 2
    r = min(r, min(w, h) / 2)
    for p in on:
        qx, qy = min(max(p.x, x0 + r), x1 - r), min(max(p.y, y0 + r), y1 - r)
        if abs(math.hypot(p.x - qx, p.y - qy) - r) > 0.25:
            return None
    return ('roundRect', x0, y0, w, h, min(50000, int(round(r / min(w, h) * 100000))))


def line_xml(d):
    col, w = rgb_hex(d.get('color')), d.get('width') or 0
    if not col or w <= 0 or 's' not in (d.get('type') or '') or (d.get('stroke_opacity') or 0) <= 0:
        return '<a:ln><a:noFill/></a:ln>'
    cap = d.get('lineCap')
    cap = cap[0] if isinstance(cap, (tuple, list)) else cap
    cap = {0: 'flat', 1: 'rnd', 2: 'sq'}.get(int(cap or 0), 'flat')
    join = int(d.get('lineJoin') or 0)
    join = {0: '<a:miter lim="800000"/>', 1: '<a:round/>', 2: '<a:bevel/>'}.get(join, '<a:miter lim="800000"/>')
    dash = '<a:prstDash val="solid"/>'
    nums = [float(v) for v in re.findall(r'[\d.]+', (d.get('dashes') or '').split(']')[0])]
    if nums and any(nums):
        if len(nums) % 2:
            nums = nums * 2
        dash = '<a:custDash>%s</a:custDash>' % ''.join(
            '<a:ds d="%d" sp="%d"/>' % (max(1, round(nums[i] / w * 100000)), max(1, round(nums[i + 1] / w * 100000)))
            for i in range(0, len(nums), 2))
    return '<a:ln w="%d" cap="%s">%s%s%s</a:ln>' % (emu(w), cap, solid(col, d.get('stroke_opacity') or 1), dash, join)


def fill_xml(d):
    col = rgb_hex(d.get('fill'))
    if not col or 'f' not in (d.get('type') or '') or (d.get('fill_opacity') or 0) <= 0:
        return '<a:noFill/>'
    return solid(col, d.get('fill_opacity') or 1)


def shadow_xml(sh):
    if not sh:
        return ''
    return ('<a:effectLst><a:outerShdw blurRad="%d" dist="%d" dir="%d" algn="ctr" rotWithShape="0">%s'
            '</a:outerShdw></a:effectLst>' % (emu(sh['blur']), emu(sh['dist']), sh['dir'],
                                             solid(sh['col'], sh['alpha'])[len('<a:solidFill>'):-len('</a:solidFill>')]))


def shape_xml(sid, name, drawings, geom_override=None, fill_override=None):
    """One p:sp for one or more drawings that share a style."""
    d0 = drawings[0]
    allsubs = [subpaths(d) for d in drawings]
    pts = [p for subs in allsubs for s in subs for o in s['ops'] for p in o[1:]]
    if not pts:
        return None
    x0, y0 = min(p.x for p in pts), min(p.y for p in pts)
    x1, y1 = max(p.x for p in pts), max(p.y for p in pts)
    W, H = max(1, emu(x1 - x0)), max(1, emu(y1 - y0))
    prst = detect_prst(allsubs[0]) if len(drawings) == 1 else None
    if prst:
        kind, px, py, pw, ph, adj = prst
        av = '<a:avLst><a:gd name="adj" fmla="val %d"/></a:avLst>' % adj if kind == 'roundRect' else '<a:avLst/>'
        geom = '<a:prstGeom prst="%s">%s</a:prstGeom>' % (kind, av)
        off, ext = (emu(px), emu(py)), (max(1, emu(pw)), max(1, emu(ph)))
    else:
        def pt(p):
            return '<a:pt x="%d" y="%d"/>' % (emu(p.x - x0), emu(p.y - y0))
        paths = []
        filled = 'f' in (d0.get('type') or '')
        for d, subs in zip(drawings, allsubs):
            groups = [subs]
            if filled and not d.get('even_odd') and len(subs) > 1:
                signs = {signed_area(s) > 0 for s in subs}
                if len(signs) == 1:
                    groups = [[s] for s in subs]
            for g in groups:
                body = ''
                for s in g:
                    for o in s['ops']:
                        if o[0] == 'M':
                            body += '<a:moveTo>%s</a:moveTo>' % pt(o[1])
                        elif o[0] == 'L':
                            body += '<a:lnTo>%s</a:lnTo>' % pt(o[1])
                        else:
                            body += '<a:cubicBezTo>%s%s%s</a:cubicBezTo>' % (pt(o[1]), pt(o[2]), pt(o[3]))
                    if s['closed']:
                        body += '<a:close/>'
                attrs = '' if filled else ' fill="none"'
                paths.append('<a:path w="%d" h="%d"%s>%s</a:path>' % (W, H, attrs, body))
        geom = ('<a:custGeom><a:avLst/><a:gdLst/><a:ahLst/><a:cxnLst/><a:rect l="0" t="0" r="r" b="b"/>'
                '<a:pathLst>%s</a:pathLst></a:custGeom>' % ''.join(paths))
        off, ext = (emu(x0), emu(y0)), (W, H)
    return ('<p:sp %s><p:nvSpPr><p:cNvPr id="%d" name="%s %d"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
            '<p:spPr><a:xfrm><a:off x="%d" y="%d"/><a:ext cx="%d" cy="%d"/></a:xfrm>%s%s%s%s</p:spPr></p:sp>'
            % (NS, sid, name, sid, off[0], off[1], ext[0], ext[1], geom, fill_override or fill_xml(d0),
               line_xml(d0), shadow_xml(d0.get('shadow'))))


# ------------------------------------------------------------------ text

class Font:
    """Measuring fonts: the kit's own file for each face Chrome used, Helvetica (Arial's widths) otherwise."""
    cache = {}

    @classmethod
    def get(cls, ps):
        if ps not in cls.cache:
            face = fontinfo.faces().get(ps) if ps else None
            cls.cache[ps] = fitz.Font(fontfile=str(face.path)) if face else fitz.Font('helv')
        return cls.cache[ps]


# Vertical metrics each app lays text out with, as fractions of the font size: 'hhea' (what Chrome on a Mac
# uses) or 'win' (OS/2 winAscent and winDescent, what Windows apps use). Measured with --check and --powerpoint.
METRICS = os.environ.get('PC_PPTX_METRICS', 'hhea')


def metrics(ps):
    """(ascent, descent) of a face, Arial's for text in a font the kit does not have."""
    f = fontinfo.faces().get(ps) if ps else None
    if f is None:
        return 0.905, 0.212
    return (f.win_asc, f.win_desc) if METRICS == 'win' else (f.asc, f.desc)


@lru_cache(None)
def face_font(ps):
    return fitz.Font(fontfile=str(fontinfo.faces()[ps].path))


def width_miss(sp, ps):
    """How far a kit face's advances miss a span's glyph positions, letter spacing taken out; None if
    the span is too short to tell."""
    if ps not in fontinfo.faces() or 'chars' not in sp:     # the check's text spans carry no glyph positions
        return None
    font, size = face_font(ps), sp['size']
    ch = [c for c in sp['chars'] if c[0] > 0]
    d = sorted(math.hypot(q[2][0] - c[2][0], q[2][1] - c[2][1]) / size - font.glyph_advance(c[0])
               for c, q in zip(ch, ch[1:]) if c[0] != 32)
    if len(d) < 3:
        return None
    mid = d[len(d) // 2]
    return sum(abs(x - mid) for x in d) / len(d)


def face_of(sp, page_fonts=()):
    """PostScript name of the kit face behind a span's PDF font name, or None for a font the kit lacks.

    A span's font name can lose its style ("AtkinsonHyperlegibleNext" for ...-Bold), so a name that
    the page's own font list (page_fonts, subset prefixes removed) holds in several styles is settled
    by the span's bold (16) and italic (2) flags, then among faces on the same side of 600 by the
    glyph advances."""
    fontname, flags = sp['font'], sp.get('flags') or 0
    name = fontname.split('+')[-1]
    full = [f for f in page_fonts if f == name or f.startswith(name + '-')]
    if not full and len(name) >= 20:                     # cut short inside the style ("...-SemiBol"): the
        full = [f for f in page_fonts if f.startswith(name)]     # page's own font list has the whole name
    if len(full) > 1:
        def weight(f):
            face = fontinfo.faces().get(f)
            return face.weight if face else (700 if 'Bold' in f else 400)
        aim = 700 if flags & 16 else 400
        same = [f for f in full if bool(flags & 2) == ('Italic' in f)] or full
        same = [f for f in same if (weight(f) >= 600) == bool(flags & 16)] or same
        miss = [(width_miss(sp, f), f) for f in same] if len(same) > 1 else []
        if miss and all(m is not None for m, _ in miss):     # SemiBold or Bold: the advances tell
            full = [min(miss, key=lambda m: (round(m[0], 3), abs(weight(m[1]) - aim)))[1]]
        else:                                                # tabular figures can match in every weight
            full = [min(same, key=lambda f: abs(weight(f) - aim))]
    f = fontinfo.lookup(full[0] if len(full) == 1 else fontname, bool(flags & 2))
    return f.psname if f else None


ROTATED = []                     # rotated spans of the slide being converted, written as rotated text boxes


def glyph_runs(trace, warn, page_fonts=()):
    """Turn Chrome's text spans into positioned runs: one per span per baseline piece."""
    glyphs = []
    for sp in trace:
        if sp.get('type') not in (0, 2) or (sp.get('opacity') or 0) <= 0:
            continue
        style = face_of(sp, page_fonts)
        if style:
            USED.add(style)
        else:
            warn('text in a font the kit does not have (%s) kept as Arial' % sp['font'].split('+')[-1])
        size = sp['size']
        col = rgb_hex(sp['color'])
        dx, dy = sp['dir']
        if abs(dy) > 1e-3 or dx < 0:
            chars = [c for c in sp['chars'] if c[0] > 0]
            if chars:
                ROTATED.append({'text': ''.join(chr(c[0]) for c in chars), 'origin': chars[0][2],
                                'end': chars[-1][2], 'last': chr(chars[-1][0]), 'dir': (dx, dy), 'style': style,
                                'size': size, 'col': col, 'op': sp.get('opacity', 1), 'seq': sp['seqno']})
            continue
        for (uc, gid, (ox, oy), bb) in sp['chars']:
            ch = chr(uc) if uc > 0 else ''
            if not ch:
                continue
            last = glyphs[-1] if glyphs else None
            if last and last['seq'] == sp['seqno'] and (gid < 0 or bb[2] - bb[0] < 1e-3) \
                    and abs(last['x'] - ox) < 0.05 and abs(last['y'] - oy) < 0.05:
                last['ch'] += ch              # the second letter of a ligature (fi, ff) shares the first's glyph
                continue
            glyphs.append({'ch': ch, 'x': ox, 'y': oy, 'x1': bb[2], 'style': style, 'size': size, 'col': col,
                           'op': sp.get('opacity', 1), 'seq': sp['seqno'], 'sw': sp.get('spacewidth') or size * 0.25})
    return glyphs


def natural(style, ch, size):
    try:
        return Font.get(style).text_length(ch, fontsize=size)
    except Exception:
        return size * 0.5


def build_lines(glyphs):
    """Group glyphs into baseline rows, split rows into segments at gaps that are not a word space."""
    glyphs.sort(key=lambda g: (g['y'], g['x']))
    rows, cur = [], []
    for g in glyphs:
        if cur and abs(g['y'] - cur[-1]['y']) > 0.4:
            rows.append(cur)
            cur = []
        cur.append(g)
    if cur:
        rows.append(cur)
    lines = []
    for row in rows:
        row.sort(key=lambda g: g['x'])
        known = row_tracking(row)
        seg = [row[0]]
        for g in row[1:]:
            p = seg[-1]
            size = max(p['size'], g['size'])
            gap = g['x'] - p['x1']
            track = tracking_of(seg) or known.get((p['style'], p['size']), 0.0)
            if p['ch'] == ' ' or g['ch'] == ' ':
                ok = -0.3 * size < gap - track < 0.3 * size
            else:
                ok = -0.3 * size < gap - track <= 0.12 * size or abs(gap - track - p['sw']) <= 0.35 * p['sw']
                if ok and gap - track > 0.12 * size:
                    seg.append(dict(p, ch=' ', x=p['x1'], x1=g['x'], virtual=True))
            if not ok:
                lines.append(seg)
                seg = []
            seg.append(g)
        lines.append(seg)
    return [make_line(s) for s in lines if any(g['ch'].strip() for g in s)]


def row_tracking(row):
    """{(face, size): median letter-spacing} over a whole row, for the first letters of a segment, before
    tracking_of has enough of them (a short letter-spaced label would otherwise split at its first wide gap)."""
    by = {}
    for a, b in zip(row, row[1:]):
        if a['ch'] != ' ' and b['ch'] != ' ' and a['style'] == b['style'] and a['size'] == b['size']:
            d = (b['x'] - a['x']) - natural(a['style'], a['ch'], a['size'])
            if abs(d) < 0.3 * a['size']:
                by.setdefault((a['style'], a['size']), []).append(d)
    return {k: sorted(v)[len(v) // 2] for k, v in by.items() if len(v) >= 3}


def tracking_of(glyphs):
    """Median extra advance between glyphs of one run (CSS letter-spacing), in points."""
    ds = []
    for a, b in zip(glyphs, glyphs[1:]):
        if a['ch'] != ' ' and b['ch'] != ' ' and not a.get('virtual') and a['style'] == b['style'] \
                and a['size'] == b['size']:
            ds.append((b['x'] - a['x']) - natural(a['style'], a['ch'], a['size']))
    if len(ds) < 3:
        return 0.0
    ds.sort()
    return ds[len(ds) // 2]


def odd_advances(seg):
    """Extra spacing (points) after each glyph of a line, and how far to move the line's start, for glyphs
    Chrome placed far from the font's own advance: proportional figures, which PowerPoint cannot turn on
    (Newsreader's "1" is 0.2 em narrower than its default), or tabular ones in tables. Each such glyph keeps
    its ink centred as Chrome drew it: half the difference goes before it, half after. Most glyphs get 0.
    LibreOffice drops negative spacing after the last letter of a run (PowerPoint applies it; tested 07 Oct
    2026), so its --check render still shows such a figure at the font's default width."""
    extra, lead = [0.0] * len(seg), 0.0
    for k, g in enumerate(seg[:-1]):
        nxt = seg[k + 1]
        if g.get('virtual') or nxt.get('virtual'):
            continue
        d = (nxt['x'] - g['x']) - natural(g['style'], g['ch'], g['size']) - g['track']
        if abs(d) > 0.08 * g['size']:
            if k:
                extra[k - 1] += d / 2
            else:
                lead = d / 2
            extra[k] += d / 2
    return extra, lead


def make_line(seg):
    while seg and not seg[-1]['ch'].strip():
        seg = seg[:-1]
    while seg and not seg[0]['ch'].strip():
        seg = seg[1:]
    runs, cur = [], []
    for g in seg:
        key = (g['style'], g['size'], g['col'], round(g['op'], 3))
        if cur and key != cur[0]['key']:
            runs.append(cur)
            cur = []
        cur.append(dict(g, key=key))
    if cur:
        runs.append(cur)
    for r in runs:
        t = tracking_of(r)
        for g in r:
            g['track'] = t if abs(t) >= 0.05 else 0.0
    glyphs = [g for r in runs for g in r]
    extra, lead = odd_advances(glyphs)
    out, cur = [], []
    for g, e in zip(glyphs, extra):                  # a run per stretch of one face and one spacing
        g['track'] += e
        if cur and (g['key'], round(g['track'], 2)) != (cur[0]['key'], round(cur[0]['track'], 2)):
            out.append(cur)
            cur = []
        cur.append(g)
    if cur:
        out.append(cur)
    out = [{'text': ''.join(g['ch'] for g in r), 'style': r[0]['style'], 'size': r[0]['size'], 'col': r[0]['col'],
            'op': r[0]['op'], 'track': r[0]['track']} for r in out]
    size = max(g['size'] for g in seg)
    # the line's metrics are those of its widest run in its biggest size (the face that sets the line box)
    main = max((r for r in out if r['size'] == size), key=lambda r: len(r['text']))
    asc, desc = metrics(main['style'])
    x0, x1 = seg[0]['x'] + lead, seg[-1]['x1']
    return {'runs': out, 'x0': x0, 'x1': x1, 'y': seg[0]['y'], 'size': size, 'seq': min(g['seq'] for g in seg),
            'last': max(g['seq'] for g in seg), 'text': ''.join(g['ch'] for g in seg), 'asc': asc, 'desc': desc}


def logical_lines(lines, entries):
    """Sinhala and Tamil lines with their text from the page (dom_lines) in place of the PDF's. Each page line
    belongs to the PDF line on its row that it overlaps most; pieces of one row that share page lines (Chrome's
    out-of-order glyphs split a row apart) become one line again. Only rows with such script change. The new
    text goes in one run per script, each in the face Chrome used for it, so digits and punctuation keep the
    Latin face; a change of weight inside such a line is not kept."""
    if not entries:
        return lines
    def box(ln):
        return ln['x0'], ln['y'] - ln['asc'] * ln['size'], ln['x1'], ln['y'] + ln['desc'] * ln['size']
    owner = {}
    for j, e in enumerate(entries):
        best = None
        for i, ln in enumerate(lines):
            x0, y0, x1, y1 = box(ln)
            w = min(x1, e['x1']) - max(x0, e['x'])
            h = min(y1, e['y1']) - max(y0, e['y'])
            if w <= 0.3 * min(x1 - x0, e['x1'] - e['x']) or h <= 0.5 * min(y1 - y0, e['y1'] - e['y']):
                continue
            if best is None or h > best[0] + 0.5:
                best = (h, ln['y'], [i])
            elif abs(ln['y'] - best[1]) < 0.5:
                best[2].append(i)
        if best:
            for i in best[2]:
                owner.setdefault(i, set()).add(j)
    groups = []                                        # [line indices, entry indices], one per merged row piece
    for i in sorted(owner):
        hit = [g for g in groups if g[1] & owner[i] and abs(lines[g[0][0]]['y'] - lines[i]['y']) < 0.5]
        for g in hit[1:]:
            hit[0][0] += g[0]
            hit[0][1] |= g[1]
            groups.remove(g)
        if hit:
            hit[0][0].append(i)
            hit[0][1] |= owner[i]
        else:
            groups.append([[i], set(owner[i])])
    for i, ln in enumerate(lines):                     # a mark Chrome drew off the baseline (the "ර" of "ප්\u200dර")
        if i in owner or not INDIC.search(ln['text']):  # is a row of its own, which no page line chose:
            continue                                   # it joins the row piece it sits on
        cx = (ln['x0'] + ln['x1']) / 2
        for idx, ents in groups:
            near = lines[idx[0]]
            if any(lines[k]['x0'] <= cx <= lines[k]['x1'] for k in idx) and abs(ln['y'] - near['y']) < 0.6 * near['size']:
                idx.append(i)
                break
    drop, out = set(), []
    for idx, ents in groups:
        parts = [lines[i] for i in idx if i in owner]
        if not any(INDIC.search(p['text']) for p in parts):
            continue
        runs = [r for p in parts for r in p['runs'] if r['text'].strip()]
        script = next((r for r in runs if INDIC.search(r['text'])), runs[0])
        es = sorted((entries[j] for j in ents), key=lambda e: e['x'])
        latin = next((r for r in runs if not INDIC.search(r['text'])), None)
        if latin is None:                              # no Latin letter in the row: spaces still take the stack's
            style = face_named(es[0].get('f', ''), es[0].get('w', 400), es[0].get('i', False))   # first font
            if style:
                USED.add(style)
            latin = dict(script, style=style) if style else script
        size = max(p['size'] for p in parts)
        text = es[0]['t']
        for a, b in zip(es, es[1:]):
            text += (' ' if b['x'] - a['x1'] > 0.25 * size else '') + b['t']
        new, cur = [], None
        for ch in text:                                # Chrome draws spaces in the stack's first font, the Latin one
            kind = script if INDIC.search(ch) or ch in '\u200c\u200d' else latin
            if new and kind is cur:
                new[-1]['text'] += ch
            else:
                new.append({'text': ch, 'style': kind['style'], 'size': kind['size'], 'col': kind['col'],
                            'op': kind['op'], 'track': 0.0})
            cur = kind
        main = max((r for r in new if r['size'] == max(n['size'] for n in new)), key=lambda r: len(r['text']))
        asc, desc = metrics(main['style'])
        out.append(dict(parts[0], runs=new, text=text, x0=min(p['x0'] for p in parts),
                        x1=max(p['x1'] for p in parts), seq=min(p['seq'] for p in parts),
                        last=max(p['last'] for p in parts), size=size, asc=asc, desc=desc))
        drop.update(idx)
    return [ln for i, ln in enumerate(lines) if i not in drop] + out


def separated(a, b, drawings):
    """True when a shape edge lies in the gap between two lines (a divider, a card edge)."""
    top, bot = a['y'] + a['desc'] * a['size'], b['y'] - b['asc'] * b['size']
    lo, hi = max(a['x0'], b['x0']), min(a['x1'], b['x1'])
    for d in drawings:
        r = d['rect']
        if r.x1 < lo or r.x0 > hi:
            continue
        if top - 0.5 < r.y0 < bot + 0.5 or top - 0.5 < r.y1 < bot + 0.5:
            return True
    return False


def build_boxes(lines, drawings):
    lines.sort(key=lambda l: (l['y'], l['x0']))
    boxes = []
    for ln in lines:
        home = None
        for b in reversed(boxes):
            last = b['lines'][-1]
            if ln['y'] - last['y'] > 2.2 * b['size']:
                continue
            if abs(ln['size'] - b['size']) > 0.3:
                continue
            pitch = ln['y'] - last['y']
            if not (0.95 * b['size'] <= pitch <= 2.2 * b['size']):
                continue
            cx, lcx = (last['x0'] + last['x1']) / 2, (ln['x0'] + ln['x1']) / 2
            align = None
            if abs(ln['x0'] - b['lines'][0]['x0']) <= 0.8 and b['align'] in (None, 'l'):
                align = 'l'
            elif abs(lcx - cx) <= 0.8 and b['align'] in (None, 'ctr'):
                align = 'ctr'
            elif abs(ln['x1'] - last['x1']) <= 0.8 and b['align'] in (None, 'r'):
                align = 'r'
            if not align or separated(last, ln, drawings):
                continue
            home = b
            b['align'] = align
            break
        if home:
            home['lines'].append(ln)
        else:
            boxes.append({'lines': [ln], 'size': ln['size'], 'align': None, 'pitch': None})
    return [part for b in boxes for part in split_paragraphs(b)]


def split_paragraphs(box):
    """Split a column of lines into one box per paragraph or list item. Inside a paragraph the lines sit
    one line pitch apart (the smallest gap in the column), give or take a pixel where Chrome grows a line
    box for a quote mark or a fallback glyph; a clearly bigger gap starts the next paragraph. A box per
    paragraph keeps each one anchored where Chrome drew it, so an app that redoes the line spacing (Google
    Slides) cannot open seams or overlaps inside a paragraph, and no box straddles two list items."""
    lines = box['lines']
    if len(lines) < 3:
        return [box]
    gaps = [b['y'] - a['y'] for a, b in zip(lines, lines[1:])]
    pitch, tol = min(gaps), max(0.35, 0.05 * box['size'])
    parts = [[lines[0]]]
    for gap, ln in zip(gaps, lines[1:]):
        if gap > pitch + tol:
            parts.append([])
        parts[-1].append(ln)
    return [dict(box, lines=part) for part in parts]


SINHALA, TAMIL = re.compile('[඀-෿]'), re.compile('[஀-௿]')


def lang_of(text):
    return 'si-LK' if SINHALA.search(text) else 'ta-LK' if TAMIL.search(text) else 'en-GB'


def rpr(run, tag='a:rPr'):
    f = fontinfo.faces().get(run['style']) if run['style'] else None
    face = f.family if f else 'Arial'
    attrs = 'lang="%s" sz="%d" b="%d" i="%d" kern="0" dirty="0"' % (
        lang_of(run.get('text', '')), round(run['size'] * 100), bool(f and f.bold), bool(f and f.italic))
    if run.get('track'):
        attrs += ' spc="%d"' % round(run['track'] * 100)
    return '<%s %s>%s<a:latin typeface="%s"/><a:cs typeface="%s"/></%s>' % (
        tag, attrs, solid(run['col'] or '000000', run['op']), esc(face), esc(face), tag)


# Helvetica, built into PyMuPDF, has Arial's exact widths. Arial is what an app without the deck's fonts draws
# instead, and it can run wider, so each text box is made wide enough for its lines in Arial too.
FALLBACK = {False: fitz.Font('helv'), True: fitz.Font('hebo')}
# Chrome draws ff, fi, fl, ffi and ffl as single ligature glyphs. Written into PowerPoint as such, words like
# "staff" stop matching Find and spell check, and fonts without those glyphs (Arial) borrow them from another.
LIGATURES = str.maketrans({'ﬀ': 'ff', 'ﬁ': 'fi', 'ﬂ': 'fl', 'ﬃ': 'ffi', 'ﬄ': 'ffl',
                           'ﬆ': 'st'})


def heavy(style):
    f = fontinfo.faces().get(style) if style else None
    return bool(f and f.weight >= 600)


def fallback_width(line):
    return sum(FALLBACK[heavy(r['style'])].text_length(r['text'].translate(LIGATURES), fontsize=r['size'])
               + (r.get('track') or 0) * len(r['text']) for r in line['runs'])


# How the three apps lay out the lines of a text box (measured 24 Sep 2026 with test decks in
# Google Slides and LibreOffice; LibreOffice follows PowerPoint):
# - 100% line spacing is 1.2 x the font size; percentages scale it; space before a paragraph is added in points.
# - PowerPoint and LibreOffice sit each line on its descent inside its own line height, so the first baseline
#   is (height - descent) below the box top. Google Slides centres the font's ascent plus descent in each line
#   (the CSS model), so its baseline is (height - ascent - descent) / 2 + ascent below the line's top.
# Measured again 7 Oct 2026 with every kit family, at line spacings from 80% to 200%, h being
# the line height (percentage x 1.2 x size) and each next line sitting one h lower in all three apps:
# - Google Slides puts the first baseline 0.8 x min(h, 1.2 x size) below the box top, whatever the font.
# - PowerPoint puts it at A + 0.75 x (h - 1.2 x size), A being min(1.2 x ascent / (ascent + descent), 0.9) x size
#   for the font (0.9 for most; Tamil and fonts with deep descenders less), and for a line mixing such fonts
#   halfway between their lowest and highest A. Exactly 100% is an exception, never written here.
# - LibreOffice puts it at h - 0.2 x size.
# So every visual line is its own paragraph at the h that makes PowerPoint's baseline 0.96 x size down, where
# Slides puts it; any extra pitch goes in space before, and lines closer than h each get their own text box.
LINE = 1.2                       # 100% line spacing as a multiple of the font size
GS_ASC = 0.96                    # first baseline below the line top, x size, in Google Slides and (by h) PowerPoint


def ppt_ascent(line):
    vals = []
    for r in line['runs']:
        a, d = metrics(r['style'])
        vals.append(min(LINE * a / (a + d), 0.9))
    return (min(vals) + max(vals)) / 2 if vals else 0.9


def natural_h(line):
    """The line height, x size, at which PowerPoint's first baseline lands where Google Slides puts it."""
    return LINE + (GS_ASC - ppt_ascent(line)) / 0.75


def line_ppr(align, height, before):
    return ('<a:pPr algn="%s" marL="0" indent="0"><a:lnSpc><a:spcPct val="%d"/></a:lnSpc><a:spcBef><a:spcPts val="%d"/>'
            '</a:spcBef><a:spcAft><a:spcPts val="0"/></a:spcAft><a:buNone/></a:pPr>'
            % (align, round(height / LINE * 100000), round(before * 100)))


def split_tight(box):
    """One text box per run of lines of one size and line height that sit at least that height apart."""
    parts = [[box['lines'][0]]]
    for a, b in zip(box['lines'], box['lines'][1:]):
        if (b['y'] - a['y'] < natural_h(b) * b['size'] - 0.3 or abs(a['size'] - b['size']) > 0.05
                or abs(natural_h(a) - natural_h(b)) > 0.002):
            parts.append([])
        parts[-1].append(b)
    return [dict(box, lines=p) for p in parts]


def box_xml(sid, box):
    lines = box['lines']
    first = lines[0]
    align = box['align'] or 'l'
    # Google Slides and LibreOffice drop negative run spacing (spc), so each line also gets the room it takes untracked
    width = max(max(l['x1'] - l['x0'] - sum(min(0, r.get('track') or 0) * len(r['text']) for r in l['runs']),
                    fallback_width(l)) for l in lines)
    slack = width * 0.02 + 0.3 * first['size'] + 1
    if align == 'l':
        x = min(l['x0'] for l in lines)
    elif align == 'r':
        x = max(l['x1'] for l in lines) - width - slack
    else:
        x = sum((l['x0'] + l['x1']) / 2 for l in lines) / len(lines) - (width + slack) / 2
    top = first['y'] - GS_ASC * first['size']
    paras = []
    for i, ln in enumerate(lines):
        h = natural_h(ln)
        before = 0 if i == 0 else max(0, ln['y'] - lines[i - 1]['y'] - h * ln['size'])
        runs = ''.join('<a:r>%s<a:t>%s</a:t></a:r>' % (rpr(r), esc(r['text'].translate(LIGATURES))) for r in ln['runs'])
        paras.append('<a:p>%s%s%s</a:p>' % (line_ppr(align, h, before), runs, rpr(ln['runs'][-1], 'a:endParaRPr')))
    height = lines[-1]['y'] - top + (natural_h(lines[-1]) - GS_ASC) * lines[-1]['size']
    name = esc(first['text'].translate(LIGATURES)[:40])
    return ('<p:sp %s><p:nvSpPr><p:cNvPr id="%d" name="Text %d: %s"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr>'
            '<p:spPr><a:xfrm><a:off x="%d" y="%d"/><a:ext cx="%d" cy="%d"/></a:xfrm>'
            '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/></p:spPr>'
            '<p:txBody><a:bodyPr wrap="square" lIns="0" tIns="0" rIns="0" bIns="0" anchor="t" rtlCol="0">'
            '<a:noAutofit/></a:bodyPr><a:lstStyle/>%s</p:txBody></p:sp>'
            % (NS, sid, sid, name, emu(x), emu(top), emu(width + slack), emu(height), ''.join(paras)))


def rotated_xml(sid, r):
    """A rotated span (a sticker label, a vertical caption) as one text box turned to the same angle."""
    dx, dy = r['dir']
    angle = math.degrees(math.atan2(dy, dx))
    hr = natural_h({'runs': [r]})
    (ox, oy), (ex, ey) = r['origin'], r['end']
    w = math.hypot(ex - ox, ey - oy) + natural(r['style'], r['last'], r['size'])
    w = max(w, sum(FALLBACK[heavy(r['style'])].text_length(c, fontsize=r['size']) for c in r['text']))
    w += 0.3 * r['size'] + 1
    h = hr * r['size']
    lx, ly = w / 2, (hr / 2 - GS_ASC) * r['size']         # box centre, relative to the first baseline origin
    cx, cy = ox + lx * dx - ly * dy, oy + lx * dy + ly * dx
    run = dict(r, track=0)
    para = '<a:p>%s<a:r>%s<a:t>%s</a:t></a:r>%s</a:p>' % (
        line_ppr('l', hr, 0), rpr(run), esc(r['text'].translate(LIGATURES)), rpr(run, 'a:endParaRPr'))
    return ('<p:sp %s><p:nvSpPr><p:cNvPr id="%d" name="Text %d: %s"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr>'
            '<p:spPr><a:xfrm rot="%d"><a:off x="%d" y="%d"/><a:ext cx="%d" cy="%d"/></a:xfrm>'
            '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/></p:spPr>'
            '<p:txBody><a:bodyPr wrap="none" lIns="0" tIns="0" rIns="0" bIns="0" anchor="t" rtlCol="0">'
            '<a:noAutofit/></a:bodyPr><a:lstStyle/>%s</p:txBody></p:sp>'
            % (NS, sid, sid, esc(r['text'][:40]), round((angle % 360) * 60000), emu(cx - w / 2), emu(cy - h / 2),
               emu(w), emu(h), para))


# ------------------------------------------------------------------ images, gradients, shadows

def block_pixmap(block):
    try:
        pix = fitz.Pixmap(block['image'])
    except Exception:
        return None
    if pix.colorspace and pix.colorspace.n not in (1, 3):
        pix = fitz.Pixmap(fitz.csRGB, pix)
    return pix


TWINS = {}


def masked_twin(doc, block):
    """An image Chrome placed through a pattern (a dashed connector drawn by a repeating gradient) reaches
    the text layer without its soft mask, so its gaps would turn black. Find the document image with the
    same pixels and take it with its mask."""
    pb = block_pixmap(block)
    if pb is None:
        return None
    key = (id(doc), pb.width, pb.height, hashlib.md5(pb.samples).hexdigest())
    if key not in TWINS:
        TWINS[key] = None
        for xr in range(1, doc.xref_length()):
            try:
                if (doc.xref_get_key(xr, 'Subtype')[1] != '/Image' or doc.xref_get_key(xr, 'SMask')[0] != 'xref'
                        or doc.xref_get_key(xr, 'Width')[1] != str(pb.width)
                        or doc.xref_get_key(xr, 'Height')[1] != str(pb.height)):
                    continue
                p = fitz.Pixmap(doc, xr)
            except Exception:
                continue
            if p.n == pb.n and p.samples == pb.samples:
                TWINS[key] = xr
                break
    return xref_pixmap(doc, TWINS[key]) if TWINS[key] else None


def xref_pixmap(doc, xref):
    """Full-resolution image with its soft mask applied."""
    try:
        pix = fitz.Pixmap(doc, xref)
        smask = doc.xref_get_key(xref, 'SMask')
        if smask[0] == 'xref':
            mask = fitz.Pixmap(doc, int(smask[1].split()[0]))
            if pix.alpha:
                pix = fitz.Pixmap(pix, 0)
            pix = fitz.Pixmap(pix, mask)
        if pix.colorspace and pix.colorspace.n not in (1, 3):
            pix = fitz.Pixmap(fitz.csRGB, pix)
        return pix
    except Exception:
        return None


def rgb_at(pix, x, y):
    c = pix.pixel(min(max(0, x), pix.width - 1), min(max(0, y), pix.height - 1))
    return (c[0],) * 3 if pix.n - pix.alpha == 1 else tuple(c[:3])


def colour_runs(pix, horizontal):
    """Flat-colour runs along the middle row or column: [[start, end, (r, g, b)], ...] in pixels."""
    n = pix.width if horizontal else pix.height
    runs = []
    for i in range(n):
        c = rgb_at(pix, i, pix.height // 2) if horizontal else rgb_at(pix, pix.width // 2, i)
        if runs and max(abs(a - b) for a, b in zip(runs[-1][2], c)) <= 8:
            runs[-1][1] = i + 1
        else:
            runs.append([i, i + 1, c])
    wide = [r for r in runs if r[1] - r[0] > 1]
    return wide or runs


def hexc(c):
    return '%02X%02X%02X' % tuple(c[:3])


def rrect_slice(x0, y0, x1, y1, r, sa, sb):
    """Items (PyMuPDF style) for the part of a rounded rectangle between x = sa and x = sb."""
    P = fitz.Point
    k = 0.5523 * r
    lr, rr = sa < x0 + r - 0.01, sb > x1 - r + 0.01
    sa, sb = (x0 if lr else sa), (x1 if rr else sb)
    pts = []
    start = P(x0 + r, y0) if lr and r else P(sa, y0)
    cur = start
    def L(q):
        nonlocal cur
        pts.append(('l', cur, q)); cur = q
    def C(a, b, q):
        nonlocal cur
        pts.append(('c', cur, a, b, q)); cur = q
    L(P(x1 - r, y0) if rr and r else P(sb, y0))
    if rr and r:
        C(P(x1 - r + k, y0), P(x1, y0 + r - k), P(x1, y0 + r)); L(P(x1, y1 - r))
        C(P(x1, y1 - r + k), P(x1 - r + k, y1), P(x1 - r, y1))
    else:
        L(P(sb, y1))
    L(P(x0 + r, y1) if lr and r else P(sa, y1))
    if lr and r:
        C(P(x0 + r - k, y1), P(x0, y1 - r + k), P(x0, y1 - r)); L(P(x0, y0 + r))
        C(P(x0, y0 + r - k), P(x0 + r - k, y0), start)
    else:
        L(start)
    return pts


def fake_drawing(items, rect, colour, opacity=1.0):
    return {'items': items, 'rect': fitz.Rect(rect), 'type': 'f', 'fill': tuple(v / 255 for v in colour),
            'fill_opacity': opacity, 'closePath': True, 'even_odd': False, 'color': None, 'width': 0}


def find_clip(clips, bbox):
    """The clip path Chrome set for a gradient: the one whose box matches the gradient's box."""
    best, err = None, 2.0
    for c in clips:
        sc = c.get('scissor')
        if sc is None:
            continue
        e = max(abs(sc.x0 - bbox.x0), abs(sc.y0 - bbox.y0), abs(sc.x1 - bbox.x1), abs(sc.y1 - bbox.y1))
        if e < err or (best is not None and abs(e - err) < 1e-6 and len(c.get('items') or []) > len(best['items'] or [])):
            best, err = c, e
    return best


# Chrome writes a CSS repeating gradient (the kit's hatching) as a function-based shading whose PostScript function
# repeats along x. MuPDF samples such a shading on a 64 x 64 grid, so a hatch more than a few stripes long comes out
# blurred, or as a smooth gradient where Chrome sized the shading to the page. Those fills are drawn from a copy of
# the PDF in which each one is a tiling pattern of plain stripes.
REPEAT_CODE = re.compile(rb'^\s*\{\s*pop\s+dup\s+truncate\s+sub\b')


def _bin(f):
    def op(st):
        b = st.pop()
        a = st.pop()
        st.append(f(a, b))
    return op


def _un(f):
    def op(st):
        st.append(f(st.pop()))
    return op


def _logic(fb, fi):
    return _bin(lambda a, b: fb(a, b) if isinstance(a, bool) else float(fi(int(a), int(b))))


def _copy(st):
    n = int(st.pop())
    if n:
        st.extend(st[-n:])


def _index(st):
    n = int(st.pop())
    st.append(st[-1 - n])


def _roll(st):
    j, n = int(st.pop()), int(st.pop())
    if n:
        j %= n
        st[-n:] = st[-j:] + st[-n:-j] if j else st[-n:]


PS_OPS = {
    'add': _bin(lambda a, b: a + b), 'sub': _bin(lambda a, b: a - b), 'mul': _bin(lambda a, b: a * b),
    'div': _bin(lambda a, b: a / b), 'idiv': _bin(lambda a, b: float(int(a / b))),
    'mod': _bin(lambda a, b: math.fmod(int(a), int(b))), 'exp': _bin(lambda a, b: a ** b),
    'atan': _bin(lambda a, b: math.degrees(math.atan2(a, b)) % 360),
    'eq': _bin(lambda a, b: a == b), 'ne': _bin(lambda a, b: a != b), 'gt': _bin(lambda a, b: a > b),
    'ge': _bin(lambda a, b: a >= b), 'lt': _bin(lambda a, b: a < b), 'le': _bin(lambda a, b: a <= b),
    'and': _logic(lambda a, b: a and b, lambda a, b: a & b), 'or': _logic(lambda a, b: a or b, lambda a, b: a | b),
    'xor': _logic(lambda a, b: a != b, lambda a, b: a ^ b),
    'bitshift': _bin(lambda a, b: float(int(a) << int(b) if b >= 0 else int(a) >> -int(b))),
    'abs': _un(abs), 'neg': _un(lambda a: -a), 'ceiling': _un(lambda a: float(math.ceil(a))),
    'floor': _un(lambda a: float(math.floor(a))), 'round': _un(lambda a: float(math.floor(a + 0.5))),
    'truncate': _un(lambda a: float(math.trunc(a))), 'cvi': _un(lambda a: float(int(a))), 'cvr': _un(float),
    'sqrt': _un(math.sqrt), 'sin': _un(lambda a: math.sin(math.radians(a))),
    'cos': _un(lambda a: math.cos(math.radians(a))), 'ln': _un(math.log), 'log': _un(math.log10),
    'not': _un(lambda a: (not a) if isinstance(a, bool) else float(~int(a))),
    'pop': lambda st: st.pop(), 'dup': lambda st: st.append(st[-1]),
    'exch': lambda st: st.extend([st.pop(), st.pop()]),
    'copy': lambda st: _copy(st), 'index': lambda st: _index(st),
    'roll': _roll, 'true': lambda st: st.append(True), 'false': lambda st: st.append(False),
}


def ps_parse(code):
    """A Type 4 (PostScript calculator) function as nested lists."""
    toks = re.findall(rb'[{}]|[^\s{}]+', code)
    stack = [[]]
    for t in toks:
        if t == b'{':
            stack.append([])
        elif t == b'}':
            proc = stack.pop()
            stack[-1].append(proc)
        else:
            try:
                stack[-1].append(float(t))
            except ValueError:
                stack[-1].append(t.decode('latin-1'))
    body = stack[0]
    return body[0] if len(body) == 1 and isinstance(body[0], list) else body


def ps_run(proc, st):
    for x in proc:
        if isinstance(x, (float, list)):
            st.append(x)
        elif x == 'if':
            p, c = st.pop(), st.pop()
            if c:
                ps_run(p, st)
        elif x == 'ifelse':
            p2, p1, c = st.pop(), st.pop(), st.pop()
            ps_run(p1 if c else p2, st)
        else:
            PS_OPS[x](st)
    return st


def stripe_runs(code, n=512):
    """[[start, end, (r, g, b)], ...] over one period, x from 0 to 1."""
    proc = ps_parse(code)
    runs = []
    for k in range(n):
        st = ps_run(proc, [(k + 0.5) / n, 0.0])
        c = tuple(int(round(min(1.0, max(0.0, float(v))) * 255)) for v in st[-3:])
        if runs and runs[-1][2] == c:
            runs[-1][1] = (k + 1) / n
        else:
            runs.append([k / n, (k + 1) / n, c])
    return runs


def nums(s):
    return [float(v) for v in re.findall(r'[-+]?(?:\d+\.?\d*|\.\d+)(?:[eE][-+]?\d+)?', s)]


def repeat_info(doc, xref):
    """(matrix, domain, stripes) when the pattern at xref is a repeating gradient as Chrome writes it, else None."""
    get = lambda k: doc.xref_get_key(xref, k)
    if get('PatternType')[1] != '2' or get('Shading/ShadingType')[1] != '1':
        return None
    if get('Shading/ColorSpace')[1] != '/DeviceRGB':
        return None
    f = get('Shading/Function')
    if f[0] != 'xref':
        return None
    fx = int(f[1].split()[0])
    if doc.xref_get_key(fx, 'FunctionType')[1] != '4':
        return None
    code = doc.xref_stream(fx) or b''
    if not REPEAT_CODE.match(code):
        return None
    m = nums(get('Matrix')[1]) if get('Matrix')[0] == 'array' else [1, 0, 0, 1, 0, 0]
    dom = nums(get('Shading/Domain')[1]) if get('Shading/Domain')[0] == 'array' else [0, 1, 0, 1]
    try:
        runs = stripe_runs(code)
    except Exception:
        return None
    return m, dom, runs


def tile_pattern(doc, xref, m, dom, runs):
    """Rewrite the pattern as plain stripes repeated a period apart (the tile overlaps its neighbours: no seams)."""
    e = 0.25
    y0, y1 = dom[2] - 1, dom[3] + 1
    parts = []
    for shift in (-1, 0, 1):
        for a, b, c in runs:
            x0, x1 = max(-e, a + shift), min(1 + e, b + shift + 0.01)
            if x1 > x0:
                parts.append('%.4f %.4f %.4f rg %.4f %.4f %.4f %.4f re f' % (
                    c[0] / 255, c[1] / 255, c[2] / 255, x0, y0, x1 - x0, y1 - y0))
    head = ('<< /Type /Pattern /PatternType 1 /PaintType 1 /TilingType 1 /BBox [ %g %.4f %g %.4f ] /XStep 1 '
            '/YStep %.4f /Matrix [ %s ] /Resources << >> >>' % (-e, y0, 1 + e, y1, y1 - y0 + 2,
                                                                 ' '.join('%.6f' % v for v in m)))
    doc.update_object(xref, head)
    doc.update_stream(xref, '\n'.join(parts).encode(), new=True)


def page_patterns(doc, page):
    """Xrefs of the patterns in the page's own resources."""
    res = doc.xref_get_key(page.xref, 'Resources')
    if res[0] == 'xref':
        pat = doc.xref_get_key(int(res[1].split()[0]), 'Pattern')
    else:
        pat = doc.xref_get_key(page.xref, 'Resources/Pattern')
    if pat[0] == 'xref':
        pat = ('dict', doc.xref_object(int(pat[1].split()[0])))
    return [int(x) for x in re.findall(r'/[^\s/<>\[\]()]+\s+(\d+)\s+0\s+R', pat[1])] if pat[0] == 'dict' else []


def patch_repeats(doc):
    """Swap every repeating gradient in doc for an equal tiling pattern, which MuPDF draws exactly. Returns
    {page number: [box]}: MuPDF logs each such fill as 'fill-shade' with the gradient's whole box, the one given here."""
    boxes, seen = {}, {}
    for page in doc:
        for xref in page_patterns(doc, page):
            if xref not in seen:
                seen[xref] = repeat_info(doc, xref)
                if seen[xref]:
                    tile_pattern(doc, xref, *seen[xref])
            if seen[xref]:
                m, dom, runs = seen[xref]
                M = fitz.Matrix(*m) * page.transformation_matrix
                pts = [fitz.Point(x, y) * M for x in dom[:2] for y in dom[2:4]]
                boxes.setdefault(page.number, []).append(
                    fitz.Rect(min(p.x for p in pts), min(p.y for p in pts), max(p.x for p in pts), max(p.y for p in pts)))
    return boxes


def shade_xml(ids, block, clips, warn, page=None, dpage=None):
    """Native shapes for a CSS gradient: hard-stop bars, smooth gradients, hatching. dpage: the page in the copy
    patch_repeats made, which draws hatching exactly."""
    pix = block_pixmap(block)
    bb = fitz.Rect(block['bbox'])
    if pix is None:
        return []
    clip = find_clip(clips, bb)
    if clip is not None and clip.get('items'):
        cd = {'items': clip['items'], 'closePath': True, 'type': 'f'}
        subs = subpaths(cd)
        prst = detect_prst(subs)
        pts = [p for sb in subs for o in sb['ops'] for p in o[1:]]
        geo = fitz.Rect(min(p.x for p in pts), min(p.y for p in pts), max(p.x for p in pts), max(p.y for p in pts))
        radius = (prst[5] / 100000 * min(geo.width, geo.height)) if prst and prst[0] == 'roundRect' else 0
        base_items = clip['items']
        if not (bb + (-1, -1, 1, 1)).contains(geo):
            # The clip is matched by its box, which the clips around it have already cut down, but its path is
            # whole (a gradient in a cropped SVG picture, or larger than its overflow: hidden box). Cut the path
            # to what Chrome drew, and to the rounded or round clips around it.
            wins = [[bb.tl, bb.tr, bb.br, bb.bl]] + [w for w in map(clip_window, clip.get('_stack', [])[:-1]) if w]
            cut = cut_fill(fake_drawing(base_items, geo, (0, 0, 0)), wins)
            if cut is None:
                return []
            geo, radius, base_items = cut['rect'], 0, cut['items']
    else:
        geo, radius = bb, 0
        base_items = [('re', fitz.Rect(bb), 1)]
    if block.get('repeat') and dpage is not None:     # a repeating gradient: hatching, kept as a sharp picture
        png = dpage.get_pixmap(matrix=fitz.Matrix(4, 4), clip=geo, alpha=True).tobytes('png')
        warn('striped fill kept as a picture')
        return [('PIC', png, fitz.Rect(geo), 1.0)]
    hr, vr = colour_runs(pix, True), colour_runs(pix, False)
    sx = bb.width / pix.width
    out = []
    if len(vr) == 1 and 1 <= len(hr) <= 3:
        # hard stops along x: the last colour is the track, the others sit on it
        out.append(shape_xml(ids.next(), 'Bar', [fake_drawing(base_items, geo, hr[-1][2])]))
        for k in range(len(hr) - 1):
            sa = geo.x0 if k == 0 else bb.x0 + (hr[k - 1][1] + hr[k][0]) / 2 * sx
            sb = bb.x0 + (hr[k][1] + hr[k + 1][0]) / 2 * sx
            if sb - sa < 0.2:
                continue
            items = rrect_slice(geo.x0, geo.y0, geo.x1, geo.y1, radius, sa, sb)
            out.append(shape_xml(ids.next(), 'Bar value', [fake_drawing(items, geo, hr[k][2])]))
        return [x for x in out if x]
    colours = {hexc(r[2]) for r in hr if r[1] - r[0] > 1}
    if len(hr) >= 6 and len(colours) <= 3 and page is not None:
        # stripes (CSS repeating gradients): PowerPoint has no equal, so use a sharp picture
        z = 4 / 3 * 3
        png = (dpage or page).get_pixmap(matrix=fitz.Matrix(z, z), clip=geo, alpha=True).tobytes('png')
        warn('striped fill kept as a picture')
        return [('PIC', png, fitz.Rect(geo), 1.0)]
    else:
        horizontal = len(hr) >= len(vr)
        n = pix.width if horizontal else pix.height
        stops = []
        for k in range(8):
            i = min(n - 1, int(round(k * (n - 1) / 7)))
            c = rgb_at(pix, i, pix.height // 2) if horizontal else rgb_at(pix, pix.width // 2, i)
            stops.append('<a:gs pos="%d"><a:srgbClr val="%s"/></a:gs>' % (round(k * 100000 / 7), hexc(c)))
        fill = ('<a:gradFill rotWithShape="1"><a:gsLst>%s</a:gsLst><a:lin ang="%d" scaled="0"/></a:gradFill>'
                % (''.join(stops), 0 if horizontal else 5400000))
    x = shape_xml(ids.next(), 'Gradient', [fake_drawing(base_items, geo, (0, 0, 0))], fill_override=fill)
    return [x] if x else []


def masked_fill_png(block, colour, opacity):
    """A soft-masked fill (Chrome's blurred shadows) as an RGBA picture."""
    mask = block_pixmap(block)
    if mask is None:
        return None
    if mask.n - mask.alpha != 1:
        mask = fitz.Pixmap(fitz.csGRAY, mask)
    rgba = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, mask.width, mask.height), True)
    src = mask.samples
    step = mask.n
    col = bytes(int(round(v * 255)) for v in colour[:3])
    buf = bytearray(mask.width * mask.height * 4)
    for i in range(mask.width * mask.height):
        buf[4 * i:4 * i + 3] = col
        buf[4 * i + 3] = int(src[i * step] * opacity)
    rgba = fitz.Pixmap(fitz.csRGB, mask.width, mask.height, bytes(buf), True)
    return rgba.tobytes('png')


def pair_images(page, bl):
    """Pair every image-like operation (image, image mask, gradient) with the image block MuPDF
    made for it, and the xref where the image has one."""
    blocks = [b for b in page.get_text('dict', flags=fitz.TEXT_PRESERVE_IMAGES)['blocks'] if b['type'] == 1]
    infos = page.get_image_info(xrefs=True)
    ops = [(i, t, fitz.Rect(r)) for i, (t, r) in enumerate(bl) if t in ('fill-image', 'fill-imgmask', 'fill-shade')]
    if len(infos) != len(blocks):
        infos = [None] * len(blocks)
    if len(ops) == len(blocks):
        return [(op, blocks[k], infos[k]) for k, op in enumerate(ops)]
    out, free = [], list(range(len(blocks)))
    for op in ops:
        if not free:
            break
        if op[1] == 'fill-shade':
            k = free[0]
        else:
            k = max(free, key=lambda j: abs(fitz.Rect(blocks[j]['bbox']) & op[2]))
        free.remove(k)
        out.append((op, blocks[k], infos[k]))
    return out


# ------------------------------------------------------------------ clips and group opacity

def flat_points(items):
    pts = []
    for it in items:
        if it[0] == 're':
            r = it[1]
            pts += [r.tl, r.tr, r.br, r.bl]
        elif it[0] == 'qu':
            q = it[1]
            pts += [q.ul, q.ur, q.lr, q.ll]
        elif it[0] == 'l':
            pts += [it[1], it[2]]
        elif it[0] == 'c':
            a, b, c, d = it[1:5]
            for t in (0, 0.25, 0.5, 0.75, 1):
                u = 1 - t
                pts.append(fitz.Point(u ** 3 * a.x + 3 * u * u * t * b.x + 3 * u * t * t * c.x + t ** 3 * d.x,
                                      u ** 3 * a.y + 3 * u * u * t * b.y + 3 * u * t * t * c.y + t ** 3 * d.y))
    return pts


def inside(pt, poly):
    x, y, hit = pt.x, pt.y, False
    for i in range(len(poly)):
        a, b = poly[i], poly[i - 1]
        if (a.y > y) != (b.y > y) and x < (b.x - a.x) * (y - a.y) / (b.y - a.y) + a.x:
            hit = not hit
    return hit


def paint_context(ext):
    """Walk PyMuPDF's extended drawing list. Returns, per path seqno, the clips in force and the
    opacity of the groups around it, plus every group's seqno span for text and images.
    Chrome draws a dashed border four times, each copy clipped to one side, and CSS opacity
    lives on a group, not on the shapes and text inside it."""
    stack, ctx, groups, last = [], {}, [], -1
    for e in ext:
        lv = e.get('level', 0)
        while stack and stack[-1][0] >= lv:
            _, kind, obj = stack.pop()
            if kind == 'group':
                obj['end'] = last
        if e['type'] == 'clip':
            if e.get('items'):
                stack.append((lv, 'clip', e))
                e['_stack'] = [obj for _, kind, obj in stack if kind == 'clip']    # for gradients, which find
                                                                                   # their clip by its box
        elif e['type'] == 'group':
            g = {'opacity': e.get('opacity') if e.get('opacity') is not None else 1.0,
                 'rect': fitz.Rect(e['rect']), 'lo': last, 'hi': None, 'end': None}
            groups.append(g)
            stack.append((lv, 'group', g))
        else:
            sq = e['seqno']
            for g in groups:
                if g['hi'] is None and g['end'] is not None:
                    g['hi'] = sq
            op, grect = 1.0, None
            for _, kind, obj in stack:
                if kind == 'group':
                    op *= obj['opacity']
                    grect = obj['rect']
            ctx[sq] = ([obj for _, kind, obj in stack if kind == 'clip'], op, grect)
            last = sq
    for g in groups:
        if g['hi'] is None:
            g['hi'] = 1 << 30
    return ctx, [g for g in groups if g['opacity'] < 0.999]


def group_opacity(groups, seq, rect):
    op = 1.0
    for g in groups:
        if g['lo'] < seq < g['hi'] and (g['rect'] + (-1, -1, 1, 1)).contains(rect):
            op *= g['opacity']
    return op


def sub_polygon(sub):
    pts, prev = [sub['ops'][0][1]], sub['ops'][0][1]
    for o in sub['ops'][1:]:
        if o[0] == 'L':
            pts.append(o[1])
        else:
            a, b, c, d = prev, o[1], o[2], o[3]
            for t in (0.25, 0.5, 0.75, 1):
                u = 1 - t
                pts.append(fitz.Point(u ** 3 * a.x + 3 * u * u * t * b.x + 3 * u * t * t * c.x + t ** 3 * d.x,
                                      u ** 3 * a.y + 3 * u * u * t * b.y + 3 * u * t * t * c.y + t ** 3 * d.y))
        prev = o[-1]
    return pts


def clip_polygons(c):
    if '_polys' not in c:
        c['_polys'] = [sub_polygon(sb) for sb in subpaths({'items': c['items'], 'closePath': True, 'type': 'f'})
                       if len(sb['ops']) > 1]
    return c['_polys']


def in_clip(pt, c):
    """Even-odd test against every subpath of a clip (Chrome clips borders with rings)."""
    return sum(inside(pt, poly) for poly in clip_polygons(c)) % 2 == 1


def flatten(sub, n=12):
    """A subpath as a polygon, each curve cut into n straight pieces."""
    pts, prev = [sub['ops'][0][1]], sub['ops'][0][1]
    for o in sub['ops'][1:]:
        if o[0] == 'L':
            pts.append(o[1])
        else:
            a, b, c, e = prev, o[1], o[2], o[3]
            for k in range(1, n + 1):
                t = k / n
                u = 1 - t
                pts.append(fitz.Point(u ** 3 * a.x + 3 * u * u * t * b.x + 3 * u * t * t * c.x + t ** 3 * e.x,
                                      u ** 3 * a.y + 3 * u * u * t * b.y + 3 * u * t * t * c.y + t ** 3 * e.y))
        prev = o[-1]
    return pts


def clip_window(c):
    """A clip that is a rectangle, rounded rectangle or ellipse (overflow: hidden, a cropped picture), as a
    convex polygon; None for any other clip shape, such as the trapezoids Chrome clips border sides with."""
    if '_win' not in c:
        subs = [sb for sb in subpaths({'items': c['items'], 'closePath': True, 'type': 'f'}) if len(sb['ops']) > 1]
        c['_win'] = flatten(subs[0], 8) if detect_prst(subs) else None
    return c['_win']


def _edges(win):
    """Each edge of a convex polygon as (a, b, length, sign), the sign making inside positive."""
    s = 1 if sum(a.x * b.y - b.x * a.y for a, b in zip(win, win[1:] + win[:1])) > 0 else -1
    return [(a, b, math.hypot(b.x - a.x, b.y - a.y), s) for a, b in zip(win[-1:] + win[:-1], win)
            if math.hypot(b.x - a.x, b.y - a.y) > 1e-6]


def cut_polygon(pts, win):
    """Sutherland-Hodgman: the part of a polygon inside the convex polygon win."""
    for a, b, _, s in _edges(win):
        side = [s * ((b.x - a.x) * (p.y - a.y) - (b.y - a.y) * (p.x - a.x)) for p in pts]
        out = []
        for k in range(len(pts)):
            cur, prev, sc, sp = pts[k], pts[k - 1], side[k], side[k - 1]
            if (sc >= 0) != (sp >= 0):
                t = sp / (sp - sc)
                out.append(fitz.Point(prev.x + (cur.x - prev.x) * t, prev.y + (cur.y - prev.y) * t))
            if sc >= 0:
                out.append(cur)
        pts = out
        if len(pts) < 3:
            return []
    return pts


def within(pts, win, tol=0.5):
    return all(s * ((b.x - a.x) * (p.y - a.y) - (b.y - a.y) * (p.x - a.x)) >= -tol * n
               for a, b, n, s in _edges(win) for p in pts)


def cut_fill(d, wins):
    """A filled path cut exactly to the convex windows wins, subpath by subpath; None if nothing is left."""
    out, changed = [], False
    for sb in subpaths(d):
        poly = flatten(sb) if len(sb['ops']) > 1 else []
        if len(poly) >= 3 and all(within(poly, w) for w in wins):
            out.append(sb)
            continue
        changed = True
        for w in wins:
            poly = cut_polygon(poly, w) if len(poly) >= 3 else []
        if len(poly) >= 3:
            out.append({'ops': [('M', poly[0])] + [('L', p) for p in poly[1:]], 'closed': True})
    return path_from(d, out) if changed else d


def image_clips(ext, seq, r):
    """The clips in force for an image that MuPDF boxes larger than what shows. Chrome prints an SVG pattern fill
    (a hatched bar) as a bitmap of the whole SVG, drawn through the filled shape set as a clip, and MuPDF gives
    the image the whole SVG's box. get_drawings lists no images, so look at the clips set after the last path
    before the image: the first that sits inside the image's box and is clearly smaller (more than 2 pt) brings
    the clips in force."""
    cands = []
    for e in ext:
        sq = e.get('seqno')
        if sq is not None:
            if sq > seq:
                break
            cands = []
        elif e['type'] == 'clip' and e.get('_stack'):
            cands.append(e)
    for c in cands:
        sc = c.get('scissor')
        if sc is not None and (r + (-1, -1, 1, 1)).contains(sc) and (r.width - sc.width > 2 or r.height - sc.height > 2):
            return c['_stack']
    return []


def is_rect_clip(c):
    subs = [sb for sb in subpaths({'items': c['items'], 'closePath': True, 'type': 'f'}) if len(sb['ops']) > 1]
    prst = detect_prst(subs)
    return bool(prst) and prst[0] == 'rect'


def clip_masked(pix, clips, region):
    """pix, which covers region, made transparent outside the clips. PyMuPDF draws the mask: a page of
    pix's size in pixels, black, with white filled through the clips."""
    w, h = pix.width, pix.height
    sx, sy = w / region.width, h / region.height

    def pt(p):
        return '%.3f %.3f' % ((p.x - region.x0) * sx, h - (p.y - region.y0) * sy)
    ops = ['0 g 0 0 %d %d re f q' % (w, h)]
    for c in clips:
        for sb in subpaths({'items': c['items'], 'closePath': True, 'type': 'f'}):
            for o in sb['ops']:
                ops.append(pt(o[1]) + (' m' if o[0] == 'M' else ' l') if o[0] in 'ML' else
                           '%s %s %s c' % (pt(o[1]), pt(o[2]), pt(o[3])))
            ops.append('h')
        ops.append('W* n' if c.get('even_odd') else 'W n')
    ops.append('1 g 0 0 %d %d re f Q' % (w, h))
    doc = fitz.open()
    page = doc.new_page(width=w, height=h)
    xref = doc.get_new_xref()
    doc.update_object(xref, '<<>>')
    doc.update_stream(xref, '\n'.join(ops).encode())
    doc.xref_set_key(page.xref, 'Contents', '%d 0 R' % xref)
    mask = page.get_pixmap(colorspace=fitz.csGRAY, alpha=False)
    if (mask.width, mask.height) != (w, h):
        return pix
    if pix.alpha:
        a = pix.samples[pix.n - 1::pix.n]
        mask = fitz.Pixmap(fitz.csGRAY, w, h, bytes(x * y // 255 for x, y in zip(a, mask.samples)), False)
        pix = fitz.Pixmap(pix, 0)
    return fitz.Pixmap(pix, mask)


def path_from(d, subs):
    """d with its path replaced by subs; None if subs is empty."""
    if not subs:
        return None
    items = []
    for sb in subs:
        prev = sb['ops'][0][1]
        for o in sb['ops'][1:]:
            items.append(('l', prev, o[1]) if o[0] == 'L' else ('c', prev, o[1], o[2], o[3]))
            prev = o[-1]
        if sb['closed'] and (abs(prev.x - sb['ops'][0][1].x) > 0.01 or abs(prev.y - sb['ops'][0][1].y) > 0.01):
            items.append(('l', prev, sb['ops'][0][1]))
    pts = [p for it in items for p in it[1:]]
    return dict(d, items=items, closePath=True,
                rect=fitz.Rect(min(p.x for p in pts), min(p.y for p in pts), max(p.x for p in pts), max(p.y for p in pts)))


def apply_clips(d, clips):
    """Clip a path to the clips in force, subpath by subpath. Chrome draws a dashed or dotted
    border four times, each copy clipped to one side, and centres each dash on the box edge
    with only its inner half visible. So sample points across each subpath, keep the subpaths
    that are really inside, and trim axis-aligned dashes to the part that shows."""
    def cuts(c):
        sc = c.get('scissor')
        if sc is None or not (sc + (-0.5, -0.5, 0.5, 0.5)).contains(d['rect']):
            return True
        # a rounded or round clip whose box holds a filled path still cuts the parts that reach its corners
        # (a band at the top of a rounded card, the border ring of a see-through pill)
        win = d.get('type') == 'f' and not is_rect_clip(c) and clip_window(c)
        return bool(win) and not all(within(flatten(sb), win) for sb in subpaths(d) if len(sb['ops']) > 1)
    live = [c for c in clips if cuts(c)]
    if not live:
        return d
    if d.get('type') == 'f' and all(clip_window(c) for c in live):
        # a filled shape under rectangular, rounded or round clips: an SVG picture cropped by its frame, or a
        # shape larger than its overflow: hidden box. Cut it exactly, since sampling keeps or trims whole pieces.
        return cut_fill(d, [clip_window(c) for c in live])
    subs = subpaths(d)
    out, changed = [], False
    for sb in subs:
        poly = sub_polygon(sb) if len(sb['ops']) > 1 else [sb['ops'][0][1]]
        box = fitz.Rect(min(p.x for p in poly), min(p.y for p in poly), max(p.x for p in poly), max(p.y for p in poly))
        n, m = 6, 4
        hits = []
        for i in range(n):
            for j in range(m):
                pt = fitz.Point(box.x0 + (i + 0.5) * box.width / n, box.y0 + (j + 0.5) * box.height / m)
                if all(in_clip(pt, c) for c in live):
                    hits.append(pt)
        frac = len(hits) / (n * m)
        axis_rect = len(poly) in (4, 5) and all(
            min(abs(p.x - box.x0), abs(p.x - box.x1)) < 0.02 and min(abs(p.y - box.y0), abs(p.y - box.y1)) < 0.02 for p in poly)
        if frac >= 0.999:
            out.append(sb)
            continue
        changed = True
        if axis_rect and frac >= 0.15:
            hx, hy = box.width / n / 2, box.height / m / 2
            r = fitz.Rect(min(p.x for p in hits) - hx, min(p.y for p in hits) - hy,
                          max(p.x for p in hits) + hx, max(p.y for p in hits) + hy)
            out.append({'ops': [('M', r.tl), ('L', r.tr), ('L', r.br), ('L', r.bl)], 'closed': True})
        elif frac >= 0.5:
            out.append(sb)
    if not changed:
        return d
    return path_from(d, out)


def clip_rrect(clips, rect, tol=1.5):
    """The rounded clip (rounded rectangle or circle) whose box matches rect, as
    (x0, y0, x1, y1, radius); a plain rectangle clip gives radius 0."""
    best = None
    for c in clips:
        sc = c.get('scissor')
        if sc is None or max(abs(a - b) for a, b in zip(sc, rect)) > tol:
            continue
        for sb in subpaths({'items': c['items'], 'closePath': True, 'type': 'f'}):
            prst = detect_prst([sb])
            if prst and prst[0] in ('roundRect', 'rect', 'ellipse'):
                _, x, y, w, h, adj = prst
                rad = min(w, h) / 2 if prst[0] == 'ellipse' else (adj or 0) / 100000 * min(w, h)
                cand = (x, y, x + w, y + h, rad)
                if max(abs(a - b) for a, b in zip(cand[:4], rect)) <= tol and (best is None or cand[4] > best[4]):
                    best = cand
    return best


def border_ring(d):
    """A dashed or dotted CSS border arrives as many small pieces around a box. Returns
    (pieces, box) when d looks like that, else None; each piece is (bbox, round, curved, points)."""
    if d['type'] != 'f' or len(d['items']) < 12:
        return None
    subs = subpaths(d)
    if len(subs) < 6:
        return None
    r = d['rect']
    if r.width < 3 or r.height < 3:
        return None
    sizes, cents = [], []
    for sb in subs:
        pts = [o[-1] for o in sb['ops']]
        b = fitz.Rect(min(p.x for p in pts), min(p.y for p in pts), max(p.x for p in pts), max(p.y for p in pts))
        if max(b.width, b.height) > 9:
            return None
        cents.append(fitz.Point((b.x0 + b.x1) / 2, (b.y0 + b.y1) / 2))
        curved = any(o[0] == 'C' for o in sb['ops'])
        # a round dot is all curves; a square dot on a circle or a rounded corner is a sector of the
        # ring, with curved ends but straight sides
        straight = any(o[0] == 'L' and abs(o[-1] - pts[i - 1]) > max(b.width, b.height) / 4
                       for i, o in enumerate(sb['ops']) if i)
        sizes.append((b, curved and not straight, curved, pts))
    if r.height <= 1.5 * max(b.height for b, *_ in sizes) or r.width <= 1.5 * max(b.width for b, *_ in sizes):
        return None                        # one row of pieces along a line (a dashed SVG line): not a box
    half = min(r.width, r.height) / 2
    best = 0
    for rad in (0, half * 0.25, half * 0.5, half * 0.75, half):
        ok = 0
        for c in cents:
            qx, qy = min(max(c.x, r.x0 + rad), r.x1 - rad), min(max(c.y, r.y0 + rad), r.y1 - rad)
            if rad > 0 and (qx != c.x or qy != c.y):
                dist = abs(math.hypot(c.x - qx, c.y - qy) - rad)
            else:
                dist = min(c.x - r.x0, r.x1 - c.x, c.y - r.y0, r.y1 - c.y)
            ok += dist <= 2.5
        best = max(best, ok / len(cents))
    if best < 0.8:
        return None                        # pieces not on one outline: not a border
    return sizes, r


def border_xml(sid, pieces, box, colour, opacity, clips):
    """One native outline for a dashed or dotted border, drawn where the CSS border shows.
    Chrome draws each dash or dot as a piece that straddles the box edge and clips it with a
    ring (the box, and the box inset by the border width). So the width comes from the ring's
    inner edge, and the dash from each piece's length along its edge: a 2 px dot arrives
    2 px long but 4.4 px across."""
    rr = clip_rrect(clips, box, tol=2.5)
    rings = [c for c in clips if len(subpaths({'items': c['items'], 'closePath': True, 'type': 'f'})) >= 2]
    t, ring_t = None, False
    for c in reversed(rings):
        subs = subpaths({'items': c['items'], 'closePath': True, 'type': 'f'})
        outs = sorted((fitz.Rect(min(o[-1].x for o in sb['ops']), min(o[-1].y for o in sb['ops']),
                                 max(o[-1].x for o in sb['ops']), max(o[-1].y for o in sb['ops'])) for sb in subs),
                      key=lambda q: -(q.width * q.height))
        if len(outs) >= 2 and outs[0].width * outs[0].height > outs[1].width * outs[1].height > 0:
            i = outs[1]
            if rr:      # the ring's outer edge runs up to half a pixel outside the box
                t = (i.x0 - rr[0] + i.y0 - rr[1] + rr[2] - i.x1 + rr[3] - i.y1) / 4
            else:
                t = (i.x0 - outs[0].x0 + i.y0 - outs[0].y0) / 2
            if t > 0:
                ring_t = bool(rr)
                break
            t = None
    shorts = sorted(min(b.width, b.height) for b, *_ in pieces)
    longs = sorted(max(b.width, b.height) for b, *_ in pieces)
    if not t:
        t = shorts[len(shorts) // 2] / 2
    x0, y0, x1, y1, rad = rr if rr else (box.x0, box.y0, box.x1, box.y1, 0)
    # each piece's length along its edge, on the straight runs between the corners
    along, arcs = [], {}
    for b, rnd, _, pts in (pieces if ring_t else ()):
        cx, cy = (b.x0 + b.x1) / 2, (b.y0 + b.y1) / 2
        for side, dist, pos, lo, hi in (('t', cy - y0, cx, x0, x1), ('b', y1 - cy, cx, x0, x1),
                                        ('l', cx - x0, cy, y0, y1), ('r', x1 - cx, cy, y0, y1)):
            if abs(dist) <= t + 1.5 and lo + rad + 1 <= pos <= hi - rad - 1:
                along.append(((b.x1 - b.x0) if side in 'tb' else (b.y1 - b.y0), rnd))
                break
        else:   # on a rounded corner (a circle is all corner): the angle it covers round the corner's centre
            ox = x0 + rad if cx <= x0 + rad else x1 - rad if cx >= x1 - rad else None
            oy = y0 + rad if cy <= y0 + rad else y1 - rad if cy >= y1 - rad else None
            if rad > t and ox is not None and oy is not None and abs(math.hypot(cx - ox, cy - oy) - (rad - t / 2)) <= t + 1.5:
                ac = math.atan2(cy - oy, cx - ox)
                ds = [(math.atan2(p.y - oy, p.x - ox) - ac + math.pi) % (2 * math.pi) - math.pi for p in pts]
                arcs.setdefault((round(ox, 1), round(oy, 1)), []).append((ac + min(ds), ac + max(ds), rnd))
    if not along and arcs:  # no straight run (a circle): measure round the middle of the line
        along = [((a1 - a0) * (rad - t / 2), rnd) for spans in arcs.values() for a0, a1, rnd in spans]
    if along:
        lens = sorted(a for a, _ in along)
        dotted = lens[len(lens) // 2] <= 1.35 * t
        round_dots = sum(rnd for _, rnd in along) > len(along) / 2
    else:       # no ring round a known box: judge by the pieces' shapes
        dotted = round_dots = (sum(curved for _, _, curved, _ in pieces) > len(pieces) / 2
                               or longs[len(longs) // 2] < 1.3 * shorts[len(shorts) // 2])
    # the pieces straddle the box edge; the visible border is the inner band of width t
    x0, y0, x1, y1 = x0 + t / 2, y0 + t / 2, x1 - t / 2, y1 - t / 2
    rad = max(0, rad - t / 2)
    w, h = x1 - x0, y1 - y0
    if w <= 0 or h <= 0:
        return None
    adj = min(50000, int(round(rad / min(w, h) * 100000))) if rad else 0
    geom = ('<a:prstGeom prst="roundRect"><a:avLst><a:gd name="adj" fmla="val %d"/></a:avLst></a:prstGeom>' % adj
            if adj else '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>')
    if dotted:
        # Chrome's dots are one border width long, one width apart: PowerPoint's "Square Dot",
        # or its "Round Dot" when Chrome drew them round (Google Slides spaces them three widths apart)
        dash, cap = '<a:prstDash val="sysDot"/>', 'rnd' if round_dots else 'flat'
    else:
        # Chrome's thin dashes run about 3 widths on, 2 off. A custom dash would copy that in PowerPoint,
        # but Google Slides draws every custom dash as dash-dot; the "Dash" preset (4 on, 3 off) looks
        # the same in both apps
        dash, cap = '<a:prstDash val="dash"/>', 'flat'
    return ('<p:sp %s><p:nvSpPr><p:cNvPr id="%d" name="Border %d"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr><p:spPr>'
            '<a:xfrm><a:off x="%d" y="%d"/><a:ext cx="%d" cy="%d"/></a:xfrm>%s<a:noFill/>'
            '<a:ln w="%d" cap="%s">%s%s<a:round/></a:ln></p:spPr></p:sp>'
            % (NS, sid, sid, emu(x0), emu(y0), max(1, emu(w)), max(1, emu(h)), geom, emu(t), cap,
               solid(rgb_hex(colour), opacity), dash))


def round_dash(d, cl):
    """A dashed or dotted border on a small circle: Chrome fills a few wedges, clips them to the ring
    (an even-odd clip of two round paths) and to one side (a triangle), four times over. Returns
    (ring box, border width, dotted, dash period in widths) when d is such a wedge set, else None."""
    if d['type'] != 'f' or len(cl) < 2:
        return None
    side, ring = cl[-1], cl[-2]
    if not side.get('items') or any(i[0] != 'l' for i in side['items']) or len(side['items']) > 4:
        return None
    if not ring.get('even_odd') or not any(i[0] == 'c' for i in ring.get('items') or []):
        return None
    boxes = []
    for sb in subpaths(dict(ring, closePath=True)):
        pts = [p for o in sb['ops'] for p in o[1:]]
        boxes.append(fitz.Rect(min(p.x for p in pts), min(p.y for p in pts), max(p.x for p in pts), max(p.y for p in pts)))
    if len(boxes) != 2:
        return None
    outer, inner = sorted(boxes, key=lambda b: -b.width)
    for c in cl[:-2]:                        # the even-odd clip's outside is a square; the circle clip
        if any(i[0] == 'c' for i in c.get('items') or []):    # before it gives the true outer edge
            pts = [p for sb in subpaths(dict(c, closePath=True)) for o in sb['ops'] for p in o[1:]]
            b = fitz.Rect(min(p.x for p in pts), min(p.y for p in pts), max(p.x for p in pts), max(p.y for p in pts))
            if abs(b.x0 - outer.x0) < 1 and abs(b.x1 - outer.x1) < 1:
                outer = b
    t = ((outer.width - inner.width) + (outer.height - inner.height)) / 4
    if t <= 0 or abs(outer.width - outer.height) > 1 or not fitz.Rect(d['rect']).contains(outer):
        return None
    wedges = subpaths(d)
    dotted = all(o[0] in ('M', 'C') for sb in wedges for o in sb['ops'][1:-1])
    period = math.pi * (outer.width - t) / max(1, len(wedges)) / t      # dash plus gap, in border widths
    return outer, t, dotted, period


def round_dash_xml(sid, box, t, dotted, period, colour, opacity):
    x0, y0, x1, y1 = box.x0 + t / 2, box.y0 + t / 2, box.x1 - t / 2, box.y1 - t / 2
    # the preset whose dash plus gap (sysDot 2 widths, sysDash 4, dash 7) is nearest Chrome's
    prst = min((('sysDot', 2), ('sysDash', 4), ('dash', 7)), key=lambda p: abs(p[1] - period))[0]
    dash, cap = ('<a:prstDash val="sysDot"/>', 'rnd') if dotted else ('<a:prstDash val="%s"/>' % prst, 'flat')
    return ('<p:sp %s><p:nvSpPr><p:cNvPr id="%d" name="Border %d"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr><p:spPr>'
            '<a:xfrm><a:off x="%d" y="%d"/><a:ext cx="%d" cy="%d"/></a:xfrm><a:prstGeom prst="ellipse"><a:avLst/>'
            '</a:prstGeom><a:noFill/><a:ln w="%d" cap="%s">%s%s<a:round/></a:ln></p:spPr></p:sp>'
            % (NS, sid, sid, emu(x0), emu(y0), max(1, emu(x1 - x0)), max(1, emu(y1 - y0)), emu(t), cap,
               solid(rgb_hex(colour), opacity), dash))


def shadow_for(sh, card):
    """Outer-shadow settings that reproduce a blurred CSS box-shadow `sh` behind `card`."""
    r, cr = sh['rect'], card['rect']
    left, right = cr.x0 - r.x0, r.x1 - cr.x1
    top, bottom = cr.y0 - r.y0, r.y1 - cr.y1
    dx, dy = (right - left) / 2, (bottom - top) / 2
    return {'blur': max(0.5, (left + right) / 2 / 1.5), 'dist': math.hypot(dx, dy),
            'dir': int(round((math.degrees(math.atan2(dy, dx)) % 360) * 60000)),
            'col': rgb_hex(sh.get('fill')), 'alpha': sh.get('fill_opacity') or 0.1}


def find_shadows(drawings, image_rects):
    """Chrome draws a blurred box-shadow as a translucent box shaped by a soft mask, which the
    PDF reader does not report, just before the element it belongs to. The box is its own group,
    or follows an image of the mask with the same bounds. Returns {shadow seqno: card or None};
    with no card to carry a native shadow, the box is dropped rather than drawn flat."""
    by_seq = {d['seqno']: d for d in drawings}
    out = {}
    for d in drawings:
        if d['type'] != 'f' or (d.get('fill_opacity') or 1) > 0.5 or len(d['items']) != 1 or d['items'][0][0] != 're':
            continue
        r = d['rect']
        masked = any(max(abs(a - b) for a, b in zip(q, r)) < 0.5 for q in image_rects) or \
            (d.get('group_rect') is not None and max(abs(a - b) for a, b in zip(d['group_rect'], r)) < 0.5)
        card = None
        for j in range(d['seqno'] + 1, d['seqno'] + 24):
            c = by_seq.get(j)
            if c is None or c['type'] != 'f' or (c.get('fill_opacity') or 1) < 0.9:
                continue
            cr = c['rect']
            m = (cr.x0 - r.x0, r.x1 - cr.x1, cr.y0 - r.y0, r.y1 - cr.y1)
            if min(m) > 0.5 and max(m) < 48 and abs(m[0] - m[1]) <= 0.35 * max(m[0], m[1]) + 0.5:
                card = c
                break
        if card is not None or masked:
            out[d['seqno']] = card
    return out


# ------------------------------------------------------------------ one slide

def merge_rects(rects, pad=2.0):
    """Rectangles that touch or overlap (within pad), merged into one each."""
    out = []
    for r in rects:
        r = fitz.Rect(r)
        i = 0
        while i < len(out):
            if (out[i] + (-pad, -pad, pad, pad)).intersects(r):
                r |= out.pop(i)
                i = 0
            else:
                i += 1
        out.append(r)
    return out


def mask_pictures(mpage, pr):
    """[(png, rect)] for the masked shapes on a slide, cut from the masks print with transparency."""
    rects = [fitz.Rect(x['rect']) for x in mpage.get_drawings() if x.get('rect') is not None]
    rects += [fitz.Rect(b['bbox']) for b in mpage.get_image_info()]
    rects += [fitz.Rect(b['bbox']) for b in mpage.get_text('dict')['blocks']]
    rects = [r & pr for r in rects if not (r & pr).is_empty]
    if not rects:
        return []
    pix = mpage.get_pixmap(matrix=fitz.Matrix(ART_ZOOM, ART_ZOOM), alpha=True)
    out = []
    for r in merge_rects(rects):
        r = (r + (-2, -2, 2, 2)) & pr
        ir = fitz.IRect(int(r.x0 * ART_ZOOM), int(r.y0 * ART_ZOOM), int(math.ceil(r.x1 * ART_ZOOM)),
                        int(math.ceil(r.y1 * ART_ZOOM))) & fitz.IRect(0, 0, pix.width, pix.height)
        if ir.is_empty:
            continue
        sub = fitz.Pixmap(pix.colorspace, ir, True)
        sub.copy(pix, ir)
        sub.set_origin(0, 0)
        if not any(sub.samples[3::4]):
            continue
        out.append((sub.tobytes('png'), fitz.Rect(ir.x0 / ART_ZOOM, ir.y0 / ART_ZOOM, ir.x1 / ART_ZOOM, ir.y1 / ART_ZOOM)))
    return out


def text_rect(box):
    ls = box['lines']
    return fitz.Rect(min(l['x0'] for l in ls), min(l['y'] - l['asc'] * l['size'] for l in ls),
                     max(l['x1'] for l in ls), max(l['y'] + l['desc'] * l['size'] for l in ls))


def convert_slide(prs, slide, cdoc, cpage, apage, notes, warn, dpage=None, repeats=(), mpage=None, dom=()):
    spTree = slide.shapes._spTree
    ids = Ids(max([int(e.get('id')) for e in spTree.iter(qn('p:cNvPr'))] + [1]))
    W, H = prs.slide_width, prs.slide_height
    pr = cpage.rect

    # Artwork: one picture at the back, or a plain background colour.
    if apage is not None:
        pix = apage.get_pixmap(matrix=fitz.Matrix(ART_ZOOM, ART_ZOOM), alpha=False)
        if pix.samples == pix.samples[:pix.n] * (pix.width * pix.height):     # pix.is_unicolor reads pixel by pixel
            c = pix.pixel(0, 0)
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = RGBColor(*c[:3])
        else:
            data = pix.tobytes('png')
            if len(data) > 300000:                    # textured or photographic art (Chalk's grain): a JPEG at
                jpg = apage.get_pixmap(matrix=fitz.Matrix(4 / 3, 4 / 3), alpha=False).tobytes(   # CSS pixels,
                    'jpeg', jpg_quality=85)           # where grain at twice the size cost 1.3 MB a slide
                if len(jpg) < len(data) / 2:
                    data = jpg
            pic = slide.shapes.add_picture(io.BytesIO(data), 0, 0, W, H)
            pic.name = 'Slide artwork'

    bl = cpage.get_bboxlog()
    ext = cpage.get_drawings(extended=True)
    clips = [x for x in ext if x.get('type') == 'clip']
    ctx, groups = paint_context(ext)
    drawings, borders, rounds, last_ring = [], [], [], None
    for x in ext:
        if x.get('type') not in ('f', 's', 'fs') or x.get('rect') is None:
            continue
        cl, op, grect = ctx.get(x['seqno'], ([], 1.0, None))
        rd = round_dash(x, cl)
        if rd is not None:
            key = ('round', rgb_hex(x.get('fill')), tuple(round(v, 1) for v in rd[0]))
            if last_ring is None or last_ring[0] != key or x['seqno'] - last_ring[1] > 6:
                rounds.append((x['seqno'], rd, x.get('fill'), (x.get('fill_opacity') or 1) * op))
            last_ring = (key, x['seqno'])
            continue
        ring = border_ring(x)
        if ring is not None:
            key = (rgb_hex(x.get('fill')), tuple(round(v) for v in x['rect']))
            if last_ring is None or last_ring[0] != key or x['seqno'] - last_ring[1] > 6:
                borders.append((x['seqno'], ring, x.get('fill'), (x.get('fill_opacity') or 1) * op, cl))
            last_ring = (key, x['seqno'])
            continue
        if 's' in x['type'] and re.search(r'\d', (x.get('dashes') or '').split(']')[0]):
            # a dashed or dotted stroke: Chrome repeats it once per side, each copy clipped to
            # that side; keep one whole copy and let PowerPoint's dash pattern draw it
            key = ('dash', tuple(round(v, 1) for v in x['rect']), len(x['items']), rgb_hex(x.get('color')))
            if last_ring is not None and last_ring[0] == key and x['seqno'] - last_ring[1] <= 6:
                last_ring = (key, x['seqno'])
                continue
            last_ring = (key, x['seqno'])
            if op < 0.999:
                x = dict(x, stroke_opacity=(x.get('stroke_opacity') or 1) * op)
            drawings.append(x)
            continue
        if x['type'] == 'f' and len(x['items']) == 1 and x['items'][0][0] == 're':
            rr = clip_rrect(cl, x['rect'])
            if rr and rr[4] > 0.3:
                x0, y0, x1, y1, rad = rr
                x = dict(x, items=rrect_slice(x0, y0, x1, y1, rad, x0, x1), rect=fitz.Rect(x0, y0, x1, y1))
                cl = [c for c in cl if c.get('scissor') is None or
                      max(abs(a - b) for a, b in zip(c['scissor'], x['rect'])) > 1.5]
        x = apply_clips(x, cl)
        if x is not None and grect is not None:
            x['group_rect'] = grect
        if x is None:
            continue
        if op < 0.999:
            x = dict(x, fill_opacity=(x.get('fill_opacity') or 1) * op, stroke_opacity=(x.get('stroke_opacity') or 1) * op)
        drawings.append(x)
    by_seq = {d['seqno']: d for d in drawings}
    page_area = pr.width * pr.height
    items, keep = [], []
    for d in drawings:
        r = d['rect']
        if r.width * r.height > 0.98 * page_area and rgb_hex(d.get('fill')) in ('FFFFFF', None) \
                and 's' not in d['type']:
            continue                               # the page's own white background
        if (r & pr).is_empty and r.width > 0 and r.height > 0:
            continue
        keep.append(d)

    img_rects = [fitz.Rect(r) for t, r in bl if t in ('fill-image', 'fill-imgmask')]
    shadows = find_shadows(drawings, img_rects)
    consumed = set(shadows)
    for sq, card in shadows.items():
        if card is not None:
            card['shadow'] = shadow_for(by_seq[sq], card)
    for (i, t, r), block, info in pair_images(cpage, bl):
        if t == 'fill-shade':
            block['repeat'] = any(max(abs(a - b) for a, b in zip(r, q)) < 1.5 for q in repeats)
            items.append((i, 'shade', block))
            continue
        nxt = by_seq.get(i + 1)
        if nxt is not None and nxt['seqno'] in shadows and max(abs(a - b) for a, b in zip(nxt['rect'], r)) < 0.3:
            continue                                   # the soft mask of a box-shadow
        items.append((i, 'image', (block, info, r, group_opacity(groups, i, fitz.Rect(block['bbox'])))))

    for d in keep:
        if d['seqno'] not in consumed:
            items.append((d['seqno'], 'path', d))
    for sq, (pieces, box), colour, op, cl in borders:
        items.append((sq, 'border', (pieces, box, colour, op, cl)))
    for sq, rd, colour, op in rounds:
        items.append((sq, 'round', rd + (colour, op)))
    trace = cpage.get_texttrace()
    if groups:
        for sp in trace:
            op = group_opacity(groups, sp['seqno'], fitz.Rect(sp['bbox']))
            if op < 0.999:
                sp['opacity'] = (sp.get('opacity') or 1) * op
    ROTATED.clear()
    glyphs = glyph_runs(trace, warn, [f[3].split('+')[-1] for f in cpage.get_fonts()])
    fills = [(sq, fitz.Rect(o['rect'] if k == 'path' else o[1] if k == 'border' else o['bbox']))
             for sq, k, o in items if k in ('path', 'border', 'shade')]

    def under_line(r, ln):
        return r.x0 < ln['x1'] and r.x1 > ln['x0'] and \
            r.y0 < ln['y'] + ln['desc'] * ln['size'] and r.y1 > ln['y'] - ln['asc'] * ln['size']

    for b in build_boxes(logical_lines(build_lines(glyphs), dom), keep):
        first, last = min(l['seq'] for l in b['lines']), max(l['last'] for l in b['lines'])
        # A fill painted partway through the box's text (an inline code chip, a highlight, a bar under its own
        # value) belongs behind that text, so the box goes in at its last glyph, above the fill
        mid = any(first < sq < last and any(under_line(r, ln) for ln in b['lines']) for sq, r in fills)
        items.append((last if mid else first, 'text', b))
    for r in ROTATED:
        items.append((r['seq'], 'rot', r))

    items.sort(key=lambda it: it[0])
    # masked shapes: above the shapes they sit on, below the first text that overlaps them
    for png, r in (mask_pictures(mpage, pr) if mpage is not None else []):
        at = next((k for k, it in enumerate(items) if it[1] == 'text' and text_rect(it[2]).intersects(r)), len(items))
        items.insert(at, (None, 'png', (png, r, 1.0)))
    xml, group = [], []

    def flush():
        if group:
            x = shape_xml(ids.next(), 'Shape' if len(group) == 1 else 'Icon', group)
            if x:
                xml.append(x)
            group.clear()

    for seq, kind, obj in items:
        if kind == 'path':
            r = obj['rect']
            small = r.width <= 64 and r.height <= 64 and not obj.get('shadow')
            if group and small and style_key(obj) == style_key(group[0]):
                u = fitz.Rect(group[0]['rect'])
                for g in group[1:]:
                    u |= g['rect']
                u |= r
                if u.width <= 80 and u.height <= 80:
                    group.append(obj)
                    continue
            flush()
            if small:
                group.append(obj)
            else:
                x = shape_xml(ids.next(), 'Shape', [obj])
                if x:
                    xml.append(x)
            continue
        flush()
        if kind == 'text':
            for part in split_tight(obj):
                xml.append(box_xml(ids.next(), part))
        elif kind == 'rot':
            xml.append(rotated_xml(ids.next(), obj))
        elif kind == 'shade':
            xml.extend(shade_xml(ids, obj, clips, warn, cpage, dpage))
        elif kind == 'round':
            xml.append(round_dash_xml(ids.next(), *obj))
        elif kind == 'border':
            x = border_xml(ids.next(), *obj)
            if x:
                xml.append(x)
        elif kind == 'png':
            xml.append(('PIC',) + obj)
        elif kind == 'image':
            block, info, r, gop = obj
            pix = xref_pixmap(cdoc, info['xref']) if info and info.get('xref') else None
            if pix is None and info and info.get('has-mask'):
                pix = masked_twin(cdoc, block)
            pix = pix or block_pixmap(block)
            place = fitz.Rect(block['bbox'])
            vis = place & r & pr if not r.is_infinite else place & pr
            cuts = [] if r.is_infinite else image_clips(ext, seq, r)
            for c in cuts:
                if c.get('scissor') is not None:
                    vis &= c['scissor']
            cuts = [c for c in cuts if not is_rect_clip(c)]
            if pix is None or vis.is_empty:
                continue
            if abs(vis.x0 - place.x0) > 0.3 or abs(vis.y0 - place.y0) > 0.3 or \
                    abs(vis.x1 - place.x1) > 0.3 or abs(vis.y1 - place.y1) > 0.3:
                sx, sy = pix.width / place.width, pix.height / place.height
                ir = fitz.IRect(int((vis.x0 - place.x0) * sx), int((vis.y0 - place.y0) * sy),
                                int(math.ceil((vis.x1 - place.x0) * sx)), int(math.ceil((vis.y1 - place.y0) * sy)))
                ir &= fitz.IRect(0, 0, pix.width, pix.height)
                if ir.is_empty:
                    continue
                sub = fitz.Pixmap(pix.colorspace, ir, pix.alpha)
                sub.copy(pix, ir)
                sub.set_origin(0, 0)
                pix = sub
            if cuts:
                pix = clip_masked(pix, cuts, vis)
            xml.append(('PIC', pix.tobytes('png'), vis, gop))
    flush()

    for x in xml:
        if isinstance(x, tuple):
            png, vis = x[1], x[2]
            pic = slide.shapes.add_picture(io.BytesIO(png), emu(vis.x0), emu(vis.y0),
                                           max(1, emu(vis.width)), max(1, emu(vis.height)))
            pic.name = 'Picture %d' % pic.shape_id
            if len(x) > 3 and x[3] < 0.999:
                blip = pic._element.find('.//' + qn('a:blip'))
                blip.append(parse_xml('<a:alphaModFix %s amt="%d"/>' % (NS, round(x[3] * 100000))))
            ids.n = max(ids.n, pic.shape_id)
        else:
            spTree.append(parse_xml(x))
    if notes:
        slide.notes_slide.notes_text_frame.text = notes


# ------------------------------------------------------------------ the presentation

TOKENS_JS = r'''<script>addEventListener("load",function(){setTimeout(function(){
var d=document.querySelector(".pc-deck"),cs=getComputedStyle(d),o={theme:d.dataset.theme,mode:d.dataset.mode,
palette:d.dataset.palette,format:d.dataset.format,title:document.title};
["bg","bg-alt","ink","ink-2","accent","accent-2","c1","c2","c3","c4","c5","c6","font-display","font-body","font-mono"]
.forEach(function(k){o[k]=cs.getPropertyValue("--"+k).trim()});
var c=document.createElement("canvas").getContext("2d");
["bg","bg-alt","ink","ink-2","accent","accent-2","c1","c2","c3","c4","c5","c6"].forEach(function(k){
c.fillStyle="#000";c.fillStyle=o[k]||"#000";c.fillRect(0,0,1,1);var p=c.getImageData(0,0,1,1).data;
o[k]=[p[0],p[1],p[2]].map(function(v){return("0"+v.toString(16)).slice(-2)}).join("").toUpperCase()});
var pre=document.createElement("pre");pre.id="pc-tokens";pre.textContent=JSON.stringify(o);document.body.appendChild(pre)},300)});</script>'''


def deck_tokens(deck):
    """The deck's theme, mode, colours (as hex) and font stacks, read from its computed style in Chrome."""
    text = chrome.inject(Path(deck).read_text(encoding='utf-8'), body=TOKENS_JS)
    tmp = chrome.temp_beside(deck, text, 'tokens')
    try:
        dom = chrome.dump_dom(tmp, budget=5000)
    finally:
        tmp.unlink(missing_ok=True)
    m = re.search(r'<pre id="pc-tokens">(.*?)</pre>', dom, re.S)
    if not m:
        log('Note: could not read the deck\'s colours and fonts; the PowerPoint theme keeps Office defaults.')
        return {}
    import html
    import json
    return json.loads(html.unescape(m.group(1)))


# Sinhala and Tamil text cannot be read back from Chrome's PDF: it stores glyphs in the order they are drawn (a
# vowel sign that sits before its consonant comes first, "ෙකාටසක්" for "කොටසක්") and conjuncts as U+FFFD. The
# page itself holds the text in reading order, so on slides with such text this script records each visual line
# of each text node: its box in the slide's own CSS pixels and its text, for logical_lines() to put back.
INDIC = re.compile('[\u0d80-\u0dff\u0b80-\u0bff\ufffd]')
DOM_LINES_JS = r'''<script>addEventListener("load",function(){
var SCRIPT=/[඀-෿஀-௿]/;
function go(){var d=document.querySelector(".pc-deck"),W=parseFloat(getComputedStyle(d).getPropertyValue("--W"))||1920,
seg=new Intl.Segmenter(undefined,{granularity:"grapheme"}),out={};
var all=[].slice.call(d.querySelectorAll(":scope > .pc-slide"));all.forEach(function(s,idx){
if(!SCRIPT.test(s.textContent))return;
all.forEach(function(o){o.style.visibility=o===s?"visible":"hidden"});
var R=s.getBoundingClientRect(),k=R.width/W||1,lines=[],w=document.createTreeWalker(s,NodeFilter.SHOW_TEXT),n;
while((n=w.nextNode())){var p=n.parentElement;
if(!n.data.trim()||!p||p.closest(".pc-notes,noscript,.pc-ui,script,style"))continue;
var cs=getComputedStyle(p);if(cs.display==="none"||cs.visibility==="hidden")continue;
var up=cs.textTransform==="uppercase",ff=cs.fontFamily.split(",")[0].replace(/["']/g,"").trim(),
fw=parseInt(cs.fontWeight)||400,fi=cs.fontStyle==="italic",cur=null,it=seg.segment(n.data)[Symbol.iterator](),g;
while(!(g=it.next()).done){var t=g.value.segment,r=document.createRange();
r.setStart(n,g.value.index);r.setEnd(n,g.value.index+t.length);var b=r.getClientRects()[0];
if(!b||!t.trim()){if(cur)cur.t+=t;continue}
var x=(b.left-R.left)/k,y=(b.top-R.top)/k,x1=(b.right-R.left)/k,y1=(b.bottom-R.top)/k;
if(cur&&Math.abs(y-cur.y)<(cur.y1-cur.y)/2){cur.t+=t;cur.x=Math.min(cur.x,x);cur.x1=Math.max(cur.x1,x1);
cur.y1=Math.max(cur.y1,y1)}else{cur={x:x,y:y,x1:x1,y1:y1,t:t,up:up,f:ff,w:fw,i:fi};lines.push(cur)}}}
lines.forEach(function(l){l.t=l.t.replace(/\s+/g," ").trim();if(l.up)l.t=l.t.toUpperCase();delete l.up});
out[idx]=lines});
var pre=document.createElement("pre");pre.id="pc-domlines";pre.textContent=JSON.stringify(out);document.body.appendChild(pre)}
try{PCDeck.set({motion:"static"})}catch(e){}
(window.PCDeck&&PCDeck.ready?PCDeck.ready():Promise.resolve()).then(function(){setTimeout(go,300)})});</script>'''


def dom_lines(deck):
    """For each slide with Sinhala or Tamil text (by index from 0), the visual lines of its text as the page
    holds it, in PDF points: [{'x', 'y', 'x1', 'y1', 't'}]. Empty for decks without such text."""
    text = Path(deck).read_text(encoding='utf-8')
    if not INDIC.search(text):
        return {}
    tmp = chrome.temp_beside(deck, chrome.inject(text, body=DOM_LINES_JS), 'domlines')
    try:
        dom = chrome.dump_dom(tmp, budget=8000)
    finally:
        tmp.unlink(missing_ok=True)
    m = re.search(r'<pre id="pc-domlines">(.*?)</pre>', dom, re.S)
    if not m:
        log('Note: could not read the Sinhala and Tamil text from the page; it is taken from the PDF, '
            'where some letters come out of order.')
        return {}
    import html
    import json
    pt = 0.75                                          # Chrome prints 96 CSS pixels to the inch: 1 px = 0.75 pt
    return {int(k): [dict(e, x=e['x'] * pt, y=e['y'] * pt, x1=e['x1'] * pt, y1=e['y1'] * pt)
                     for e in v if e['t']] for k, v in json.loads(html.unescape(m.group(1))).items()}


def face_named(family, weight, italic):
    """The kit's face for a CSS font family, weight and style (the nearest weight it has), or None."""
    best = None
    for style, f in fontinfo.faces().items():
        if f.gfamily.lower() == family.lower() and f.italic == italic:
            if best is None or abs(f.weight - weight) < best[0]:
                best = (abs(f.weight - weight), style)
    return best and best[1]


def first_family(stack):
    m = re.match(r'\s*["\']?([^,"\']+)', stack or '')
    return m.group(1).strip() if m else ''


def scale_part(el, sx, sy, sf):
    """Scale positions, sizes and font sizes in a slide master or layout made for another slide size."""
    for e in el.iter(qn('a:off'), qn('a:chOff')):
        e.set('x', str(int(int(e.get('x')) * sx)))
        e.set('y', str(int(int(e.get('y')) * sy)))
    for e in el.iter(qn('a:ext'), qn('a:chExt')):
        if e.get('cx') is not None:
            e.set('cx', str(int(int(e.get('cx')) * sx)))
            e.set('cy', str(int(int(e.get('cy')) * sy)))
    for e in el.iter():
        if e.get('sz') is not None and e.tag in (qn('a:defRPr'), qn('a:rPr'), qn('a:endParaRPr')):
            e.set('sz', str(int(round(int(e.get('sz')) * sf / 100)) * 100))


def theme_xml(xml, tok):
    """The deck's colours and fonts in the PowerPoint theme, so shapes and text added later match it."""
    if not tok:
        return xml
    dark = tok.get('mode') == 'dark'
    lt1, dk1 = (tok['ink'], tok['bg']) if dark else (tok['bg'], tok['ink'])
    lt2, dk2 = (tok['ink-2'], tok['bg-alt']) if dark else (tok['bg-alt'], tok['ink-2'])
    colours = {'dk1': dk1, 'lt1': lt1, 'dk2': dk2, 'lt2': lt2, 'hlink': tok['accent'], 'folHlink': tok['accent-2']}
    colours.update({'accent%d' % i: tok['c%d' % i] for i in range(1, 7)})
    for k, v in colours.items():
        xml = re.sub(r'<a:%s>.*?</a:%s>' % (k, k), '<a:%s><a:srgbClr val="%s"/></a:%s>' % (k, v, k), xml, flags=re.S)
    xml = re.sub(r'(<a:clrScheme name=")[^"]*', r'\g<1>%s' % esc(tok.get('theme', 'Deck').title()), xml)
    for slot, key in (('majorFont', 'font-display'), ('minorFont', 'font-body')):
        fam = family_for_theme(first_family(tok.get(key)))
        if fam:
            xml = re.sub(r'(<a:%s>\s*<a:latin typeface=")[^"]*(")' % slot, r'\g<1>%s\2' % esc(fam), xml)
    return xml


def family_for_theme(gfamily):
    """The PowerPoint name of a Google family's regular face ("Schibsted Grotesk")."""
    for f in fontinfo.faces().values():
        if f.gfamily == gfamily and f.slot == 'regular' and f.family == f.gfamily:
            return f.family
    return gfamily


def new_presentation(width_pt, height_pt, tok):
    """An empty presentation at the deck's page size with its theme colours and fonts."""
    prs = Presentation()
    w0, h0 = prs.slide_width, prs.slide_height
    prs.slide_width, prs.slide_height = emu(width_pt), emu(height_pt)
    sx, sy = prs.slide_width / w0, prs.slide_height / h0
    sf = min(width_pt, height_pt) / (min(w0, h0) / EMU)        # type scales with the short side
    for master in prs.slide_masters:
        scale_part(master._element, sx, sy, sf)
        for layout in master.slide_layouts:
            scale_part(layout._element, sx, sy, sf)
        if tok.get('mode') == 'dark':
            cm = master._element.find(qn('p:clrMap'))
            if cm is not None:
                for k, v in (('bg1', 'dk1'), ('tx1', 'lt1'), ('bg2', 'dk2'), ('tx2', 'lt2')):
                    cm.set(k, v)
        theme = next((r.target_part for r in master.part.rels.values() if r.reltype.endswith('/theme')), None)
        if theme is not None:
            theme._blob = theme_xml(theme.blob.decode('utf-8'), tok).encode('utf-8')
    if tok.get('title'):
        prs.core_properties.title = tok['title']
    prs.core_properties.author = ''
    layout = next((l for l in prs.slide_layouts if l.name == 'Blank'), prs.slide_layouts[-1])
    return prs, layout


USED = set()                     # PostScript names of the faces the slides use


def embed_fonts(prs):
    """Embed every face the slides use, in its family's slot (regular, bold, italic, bold italic). The kit's
    fonts are all OFL with installable embedding, so PowerPoint shows them on any computer."""
    part, pres = prs.part, prs.part._element
    fams = {}
    for ps in sorted(USED):
        f = fontinfo.faces()[ps]
        fams.setdefault(f.family, {})[f.slot] = f
    if not fams:
        return []
    body = []
    for fam, slots in sorted(fams.items()):
        refs = []
        for slot in ('regular', 'bold', 'italic', 'boldItalic'):
            f = slots.get(slot)
            if not f:
                continue
            name = re.sub(r'[^a-z0-9]+', '', f.psname.lower())
            fp = Part(PackURI('/ppt/fonts/%s.fntdata' % name), 'application/x-fontdata', part.package, fontinfo.eot(f))
            refs.append('<p:%s r:id="%s"/>' % (slot, part.relate_to(fp, RT_FONT)))
        body.append('<p:embeddedFont %s><p:font typeface="%s" pitchFamily="2" charset="0"/>%s</p:embeddedFont>'
                    % (NS, esc(fam), ''.join(refs)))
    lst = pres.find(qn('p:embeddedFontLst'))
    if lst is None:
        lst = parse_xml('<p:embeddedFontLst %s/>' % NS)
        anchor = pres.find(qn('p:notesSz'))
        anchor.addnext(lst)
    for e in parse_xml('<x %s>%s</x>' % (NS, ''.join(body))):
        lst.append(e)
    pres.set('embedTrueTypeFonts', '1')
    if 'saveSubsetFonts' in pres.attrib:
        del pres.attrib['saveSubsetFonts']
    return sorted(fams)


def motion_note(src):
    """Say what the deck's motion setting (data-motion on .pc-deck) means for the PowerPoint."""
    text = Path(src).read_text(encoding='utf-8', errors='replace')
    text = re.sub(r'<!--.*?-->|<script\b.*?</script\s*>|<style\b.*?</style\s*>|<noscript\b.*?</noscript\s*>', '', text,
                  flags=re.S | re.I)
    tag = re.search(r'<[^>]*class="[^"]*\bpc-deck\b[^"]*"[^>]*>', text)
    mode = re.search(r'data-motion="(\w+)"', tag.group(0)) if tag else None
    builds = len(re.findall(r'\sdata-build(?:-each)?="', text))
    if mode and mode.group(1) == 'complex':
        warn = ('WARNING: complex motion (data-motion="complex"): not well optimised for static viewing. The HTML-only '
                'animations and interactions do not carry over; the PowerPoint shows each slide as it prints. Check '
                'every slide before presenting or sharing, and tell the user so.')
        log(warn)
        return warn
    elif builds and not (mode and mode.group(1) == 'static'):
        log('Motion: %d content builds (data-build) play in the HTML only; the PowerPoint shows every slide finished.'
            % builds)


def convert(src, out, keep=None, only=None, check=False, powerpoint=False):
    t0 = time.time()
    work = Path(keep) if keep else Path(tempfile.mkdtemp(prefix='pc-pptx-'))
    work.mkdir(parents=True, exist_ok=True)
    deck = one_file(src, work)
    pdfs = {v: work / ('%s.pdf' % v) for v in ('full', 'art', 'content', 'masks')}
    for v, pdf in pdfs.items():
        print_variant(deck, v, pdf)
    full, cdoc, adoc = fitz.open(pdfs['full']), fitz.open(pdfs['content']), fitz.open(pdfs['art'])
    mdoc = fitz.open(pdfs['masks'])
    ddoc = fitz.open(pdfs['content'])                  # a copy to draw hatching from
    repeats = patch_repeats(ddoc)
    n = full.page_count
    if not cdoc.page_count == adoc.page_count == mdoc.page_count == n:
        sys.exit('Print variants disagree on the slide count (%d, %d, %d, %d)'
                 % (n, cdoc.page_count, adoc.page_count, mdoc.page_count))
    notes = read_notes(deck)
    if len(notes) != n:
        log('Note: %d slides printed but %d .pc-slide elements found; speaker notes skipped.' % (n, len(notes)))
        notes = [''] * n
    tok = deck_tokens(deck)
    dom = dom_lines(deck)
    if dom:
        log('Sinhala or Tamil text on slide(s) %s. The kit handles these scripts at a basic level only, tested on '
            'one sample deck: before the deck goes out, look at each of these slides in PowerPoint and in Google '
            'Slides, and fix what is off.' % ', '.join(str(i + 1) for i in sorted(dom)))
    page = full[0].rect
    prs, layout = new_presentation(page.width, page.height, tok)
    USED.clear()
    warnings = {}
    for i in range(n):
        if only and (i + 1) not in only:
            continue
        slide = prs.slides.add_slide(layout)
        for ph in list(slide.placeholders):
            ph._element.getparent().remove(ph._element)
        seen = warnings.setdefault(i + 1, [])
        convert_slide(prs, slide, cdoc, cdoc[i], adoc[i], notes[i],
                      lambda m, s=seen: s.append(m) if m not in s else None, ddoc[i], repeats.get(i, ()), mdoc[i],
                      dom.get(i, ()))
    fams = embed_fonts(prs)
    prs.save(str(out))
    log('Wrote %s: %d slides, %.1f MB, %.0f s; fonts embedded: %s' % (
        out, len(prs.slides), os.path.getsize(out) / 1e6, time.time() - t0, ', '.join(fams) or 'none'))
    warn = motion_note(deck)
    for k, ws in warnings.items():
        for w in ws:
            log('  slide %d: %s' % (k, w))
    if check:
        check_pptx(out, pdfs['full'], work, only)
    if powerpoint:
        powerpoint_check(out, pdfs['full'], work, only)
    if warn and (check or powerpoint):
        log(warn)
    if not keep:
        shutil.rmtree(work, ignore_errors=True)
    return pdfs['full'] if keep else None


# ------------------------------------------------------------------ side-by-side check

def check_pptx(pptx_path, ref_pdf, work, only=None):
    soffice = find_soffice()
    if not soffice:
        log('--check skipped: LibreOffice not found (set SOFFICE_PATH).')
        return
    outdir = Path(str(pptx_path)[:-5] + ' check')
    outdir.mkdir(exist_ok=True)
    prof = Path(work) / 'lo-profile'
    subprocess.run([soffice, '-env:UserInstallation=%s' % prof.as_uri(), '--headless', '--convert-to', 'pdf',
                    '--outdir', str(work), str(pptx_path)], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                   timeout=600)
    got = Path(work) / (Path(pptx_path).stem + '.pdf')
    if not got.exists():
        log('--check: LibreOffice did not convert the file.')
        return
    ref, pp = fitz.open(ref_pdf), fitz.open(got)
    patch_repeats(ref)                                 # so the HTML side shows hatching as Chrome drew it
    fonts = sorted({f[3].split('+')[-1] for p in pp for f in p.get_fonts()})
    log('Fonts LibreOffice used: %s' % ', '.join(fonts))
    position_report(ref_pdf, got, [i for i in range(ref.page_count) if not only or (i + 1) in only],
                    'Line position test (LibreOffice)', lo=True)
    pages = [i for i in range(ref.page_count) if not only or (i + 1) in only]
    z = 960 / ref[0].rect.width
    for k, i in enumerate(pages):
        a = ref[i].get_pixmap(matrix=fitz.Matrix(z, z), alpha=False)
        zb = 960 / pp[k].rect.width
        b = pp[k].get_pixmap(matrix=fitz.Matrix(zb, zb), alpha=False)
        side = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, a.width * 2 + 12, a.height), False)
        side.clear_with(255)
        side.copy(a, a.irect)
        b.set_origin(a.width + 12, 0)
        side.copy(b, b.irect)
        side.save(str(outdir / ('slide%02d.png' % (i + 1))))
        diff = pixel_diff(a, b)
        log('  slide %2d: %4.1f%% of pixels differ' % (i + 1, diff))
    log('Side-by-side images (HTML left, LibreOffice showing the PowerPoint right): %s' % outdir)
    fallback_check(pptx_path, got, work, soffice, pages)
    layout_check(pptx_path, pages)



def family_faces():
    """{(PowerPoint family, bold, italic): Face} for the kit's faces."""
    return {(f.family, f.bold, f.italic): f for f in fontinfo.faces().values()}


RUN_RE = re.compile(r'<a:(rPr|endParaRPr|defRPr)\b([^>]*)>(.*?)</a:\1>', re.S)


def as_arial(xml):
    """Slide XML with every kit font swapped for Arial, as an app without the deck's fonts shows it: faces of
    weight 600 and up as Arial Bold (the wider case)."""
    fam = family_faces()

    def one(m):
        tag, attrs, inner = m.groups()
        t = re.search(r'typeface="([^"]*)"', inner)
        if not t:
            return m.group(0)
        b, i = ' b="1"' in attrs, ' i="1"' in attrs
        f = fam.get((esc_un(t.group(1)), b, i)) or fam.get((esc_un(t.group(1)), False, False))
        if f is None or re.search(r'lang="(?:si|ta)-', attrs):   # Arial has no Sinhala or Tamil: an app without
            return m.group(0)                                    # the deck's fonts uses its own for them
        if f.weight >= 600:
            attrs = re.sub(r' b="\d"', '', attrs) + ' b="1"'
        inner = re.sub(r'typeface="[^"]*"', 'typeface="Arial"', inner)
        return '<a:%s%s>%s</a:%s>' % (tag, attrs, inner, tag)
    return RUN_RE.sub(one, xml)


def esc_un(s):
    return s.replace('&quot;', '"').replace('&lt;', '<').replace('&gt;', '>').replace('&amp;', '&')


def fallback_check(pptx_path, got, work, soffice, pages):
    """Google Slides ignores fonts embedded in a PowerPoint; where a deck's Google font has no exact match there,
    it shows Arial. Render the file once more that way and compare text lines per slide with the first render:
    a slide that gains lines has text that re-wraps without the deck's fonts. Also count ligature characters
    (ff, fi and so on), which look right but make words fail Find and spell check in PowerPoint."""
    alt = Path(work) / 'arial-check.pptx'
    lig = 0
    with zipfile.ZipFile(pptx_path) as zi, zipfile.ZipFile(alt, 'w', zipfile.ZIP_DEFLATED) as zo:
        for n in zi.namelist():
            d = zi.read(n)
            if n.endswith('.xml'):
                x = d.decode('utf-8')
                if n.startswith('ppt/slides/slide'):
                    lig += len(re.findall('[\ufb00-\ufb06]', x))
                x = re.sub(r'<p:embeddedFontLst\b.*?</p:embeddedFontLst>', '', x, flags=re.S)
                d = as_arial(x).encode('utf-8')
            if not n.startswith('ppt/fonts/'):
                zo.writestr(n, d)
    prof = Path(work) / 'lo-profile'
    subprocess.run([soffice, '-env:UserInstallation=%s' % prof.as_uri(), '--headless', '--convert-to', 'pdf',
                    '--outdir', str(work), str(alt)], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                   timeout=600)
    alt_pdf = alt.with_suffix('.pdf')
    if not alt_pdf.exists():
        log('Arial test: LibreOffice did not convert the file.')
        return

    def lines(pdf):
        return [sum(1 for b in p.get_text('dict')['blocks'] for ln in b.get('lines', [])
                    if ''.join(sp['text'] for sp in ln['spans']).strip()) for p in fitz.open(pdf)]
    a, b = lines(got), lines(alt_pdf)
    grown = [str(pages[k] + 1) for k in range(min(len(a), len(b))) if b[k] > a[k]]
    if grown:
        log('Arial test (an app without the deck\'s fonts): text re-wraps on slide(s) %s. Compare those slides '
            'and report it.' % ', '.join(grown))
    else:
        log('Arial test (an app without the deck\'s fonts): every slide keeps its text lines.')
    if lig:
        log('Ligature characters in the slide text: %d (those words fail Find and spell check).' % lig)


def layout_check(pptx_path, pages):
    """Read every text paragraph back from the file and check its line height is natural_h() of its fonts, the
    height at which PowerPoint and Google Slides put its first baseline in the same place (see LINE). PowerPoint's
    baseline moves 0.75 x any miss; flags slides where a line could move more than 2 pt."""
    fam = family_faces()
    hits = {}
    with zipfile.ZipFile(pptx_path) as z:
        names = sorted((n for n in z.namelist() if re.match(r'ppt/slides/slide\d+\.xml$', n)),
                       key=lambda n: int(re.findall(r'\d+', n)[0]))
        for k, name in enumerate(names):
            for para in re.findall(r'<a:p>(.*?)</a:p>', z.read(name).decode('utf-8'), re.S):
                pct = re.search(r'<a:lnSpc><a:spcPct val="(\d+)"', para)
                runs, size = [], 0
                for attrs, inner in re.findall(r'<a:rPr\b([^>]*)>(.*?)</a:rPr>', para, re.S):
                    sz = re.search(r'\bsz="(\d+)"', attrs)
                    t = re.search(r'typeface="([^"]*)"', inner)
                    if not (sz and t and pct):
                        continue
                    f = fam.get((esc_un(t.group(1)), ' b="1"' in attrs, ' i="1"' in attrs))
                    runs.append({'style': f.psname if f else None})
                    size = max(size, int(sz.group(1)) / 100)
                worst = 0.0
                if runs:                                   # PowerPoint's baseline moves 0.75 x any height miss
                    h = int(pct.group(1)) / 100000 * LINE
                    worst = abs(h - natural_h({'runs': runs})) * 0.75 * size
                if worst > 2:
                    hits[pages[k] + 1] = max(hits.get(pages[k] + 1, 0), worst)
    if hits:
        log('Line height test: lines could sit differently in Google Slides on slide(s) %s (up to %.0f pt). '
            'These are lines of mixed fonts or sizes; compare them there.'
            % (', '.join(str(h) for h in sorted(hits)), max(hits.values())))
    else:
        log('Line height test: every line has the height at which PowerPoint and Google Slides agree.')
    return hits


def line_positions(pdf_path, scale=1.0, lo=False):
    """Per page: {first 30 characters of a text line: (baseline y, width)}. With lo, each baseline is moved to
    where LibreOffice puts that line of the converted file: h - 0.2 x size below the box top (see LINE), which
    is (h - GS_ASC - 0.2) x size below the HTML's."""
    out = []
    for page in fitz.open(pdf_path):
        k = scale
        pf = [f[3].split('+')[-1] for f in page.get_fonts()] if lo else ()
        lines = {}
        for b in page.get_text('dict')['blocks']:
            for ln in b.get('lines', []):
                t = ''.join(sp['text'] for sp in ln['spans']).strip()
                if len(t) > 3:
                    y = ln['spans'][0]['origin'][1]
                    if lo:
                        size = ln['spans'][0]['size']      # spans of another size can be a separate line on that baseline
                        sps = [sp for sp in ln['spans'] if sp['text'].strip() and abs(sp['size'] - size) < 0.05]
                        y += (natural_h({'runs': [{'style': face_of(sp, pf)} for sp in sps]}) - GS_ASC - 0.2) * size
                    lines.setdefault(t[:30], (y * k, (ln['bbox'][2] - ln['bbox'][0]) * k))
        out.append(lines)
    return out


def position_report(ref_pdf, got_pdf, pages, label, lo=False):
    """Compare every text line's baseline with Chrome's print (with lo, with where LibreOffice should put it);
    flag slides where one is 3 pt or more off."""
    ref, got = line_positions(ref_pdf, lo=lo), line_positions(got_pdf)
    worst, far, ratios = [], [], []
    for k, i in enumerate(pages):
        if k >= len(got):
            break
        a, b = ref[i], got[k]
        d = [abs(a[t][0] - b[t][0]) for t in a if t in b]
        ratios += [b[t][1] / a[t][1] for t in a if t in b and len(t) >= 25 and a[t][1] > 0]
        if d:
            worst.append(max(d))
            if max(d) >= 3:
                far.append('%d (%.1f pt)' % (i + 1, max(d)))
    if not worst:
        log('%s: no text to compare.' % label)
        return
    if far:
        log('%s: lines 3 pt or more from the HTML on slide(s) %s. Compare them and report it.' % (label, ', '.join(far)))
    else:
        log('%s: every line within %.1f pt of the HTML.' % (label, max(worst)))
    if ratios and sorted(ratios)[len(ratios) // 2] > 1.03:
        log('%s: text runs wider than the HTML, so the app did not use the deck\'s fonts. Report it.' % label)


# PowerPoint is sandboxed, and macOS 27 keeps tools out of its own Documents folder, so the test works in a folder
# of the kit's own that PowerPoint is granted once. The folder stays,
# since PowerPoint's access is tied to it; the files in it go after each test.
CHECK_DIR = Path(__file__).resolve().parent / 'powerpoint-check'


def powerpoint_check(pptx_path, ref_pdf, work, only=None):
    """Have PowerPoint itself (macOS, even an unlicensed view-only copy) save the file as PDF, then compare
    every line with the HTML. The first run shows PowerPoint's Grant File Access window for CHECK_DIR; once
    Grant Access is clicked there, later runs need nothing. If the test started PowerPoint, it quits it
    afterwards; if PowerPoint was already open, it closes only its own copy and leaves the rest alone."""
    app = Path('/Applications/Microsoft PowerPoint.app')
    if sys.platform != 'darwin' or not app.exists():
        log('PowerPoint test skipped: it needs PowerPoint on a Mac.')
        return
    box = CHECK_DIR
    src, pdf = box / 'check.pptx', box / 'check.pdf'
    try:
        box.mkdir(exist_ok=True)
        for f in (src, pdf):
            f.unlink(missing_ok=True)
        shutil.copy(pptx_path, src)
    except OSError as e:
        log('PowerPoint test skipped: could not write to %s (%s).' % (box, e.strerror))
        return
    was_open = subprocess.run(['pgrep', '-x', 'Microsoft PowerPoint'], capture_output=True).returncode == 0
    # the first open waits on PowerPoint's Grant File Access window, so wait for the file to be open
    script = ('with timeout of 300 seconds\n tell application "Microsoft PowerPoint"\n'
              '  open POSIX file "%s"\n  set p to missing value\n  repeat 560 times\n'
              '   try\n    set p to presentation "check.pptx"\n    exit repeat\n   end try\n   delay 0.5\n  end repeat\n'
              '  delay 1\n  save p in POSIX file "%s" as save as PDF\n  close p saving no\n end tell\nend timeout' % (src, pdf))
    err = ''
    try:
        r = subprocess.run(['osascript', '-e', script], capture_output=True, text=True, timeout=330)
        lines = (r.stderr or '').strip().splitlines()
        err = lines[-1] if lines else ''
    except Exception as e:
        err = 'PowerPoint did not respond: %s' % e
    finally:
        if not was_open:
            try:
                subprocess.run(['osascript', '-e', 'tell application "Microsoft PowerPoint" to quit'],
                               capture_output=True, timeout=60)
            except Exception:
                log('PowerPoint test: PowerPoint did not quit; close it by hand.')
    if not pdf.exists():
        log('PowerPoint test skipped: PowerPoint did not save a PDF%s. The first run asks for access to %s: click '
            'Grant Access in PowerPoint\'s window (and allow control of PowerPoint if macOS asks), then run it again.'
            % (' (%s)' % err if err else '', box))
        src.unlink(missing_ok=True)
        return
    keep = Path(work) / 'powerpoint.pdf'
    shutil.copy(pdf, keep)
    for f in (src, pdf):                               # the folder stays: PowerPoint's access is tied to it
        f.unlink(missing_ok=True)
    ref = fitz.open(ref_pdf)
    pages = [i for i in range(ref.page_count) if not only or (i + 1) in only]
    position_report(ref_pdf, keep, pages, 'PowerPoint test')


def pixel_diff(a, b):
    sa, sb = a.samples, b.samples
    n = min(len(sa), len(sb))
    step = a.n * 3
    bad = total = 0
    for j in range(0, n - a.n, step):
        total += 1
        if abs(sa[j] - sb[j]) + abs(sa[j + 1] - sb[j + 1]) + abs(sa[j + 2] - sb[j + 2]) > 90:
            bad += 1
    return 100.0 * bad / max(1, total)


def main():
    ap = argparse.ArgumentParser(description=__doc__.split('\n\n')[0])
    ap.add_argument('deck', help='the HTML deck')
    ap.add_argument('-o', '--out', help='output .pptx (default: beside the deck)')
    ap.add_argument('--check', action='store_true', help='render the result with LibreOffice and compare')
    ap.add_argument('--keep', help='keep the working files (PDFs, and the print copies of a single-file deck) in this folder')
    ap.add_argument('--powerpoint', action='store_true',
                    help='also have PowerPoint (on a Mac) save the result as PDF and compare every line')
    ap.add_argument('--slides', help='only these slides, e.g. 1,5,19 (for testing)')
    a = ap.parse_args()
    src = Path(a.deck).resolve()
    if not src.exists():
        sys.exit('No such file: %s' % src)
    out = Path(a.out).resolve() if a.out else src.with_suffix('.pptx')
    only = {int(s) for s in a.slides.split(',')} if a.slides else None
    convert(src, out, a.keep, only, a.check, a.powerpoint)


if __name__ == '__main__':
    main()
